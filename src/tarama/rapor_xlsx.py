"""Çıktıları tek bir Excel çalışma kitabında toplar → ciktilar/excel/asbest_istanbul_ikinci_tur.xlsx"""
from __future__ import annotations

import csv
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common import DOGRULANMIS, OUT, RAW, XLSX
from veri_noktalari import KADIKOY, KADIKOY_CSV, MALTEPE_MAH_2016, MALTEPE_MAH_2019

F = "Arial"
HDR = PatternFill("solid", fgColor="1F3A5F")
INPUT = Font(name=F, size=10, color="0000FF")
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)


def load(name):
    kok = DOGRULANMIS if name in ("veri_noktalari.csv", "kapsam_39.csv") else OUT
    return list(csv.DictReader(open(kok / name, encoding="utf-8-sig")))


def header(ws, row, cols, widths=None):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font = Font(name=F, size=10, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def table(ws, rows, cols, widths, start=1, link_col=None):
    header(ws, start, cols, widths)
    for i, r in enumerate(rows, start + 1):
        for j, c in enumerate(cols, 1):
            v = r.get(c, "")
            if isinstance(v, str):
                v = ILLEGAL_CHARACTERS_RE.sub("", v)
            if isinstance(v, str) and v.replace(".", "", 1).isdigit() and c not in ("sha256", "yil", "sayfa"):
                v = float(v) if "." in v else int(v)
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = BODY
            cell.alignment = Alignment(wrap_text=c in ("alinti", "sonuc", "sonraki_adim", "dizinde_gorulen",
                                                       "not", "acik_veri", "erisim"), vertical="top")
            if link_col and c == link_col and isinstance(v, str) and v.startswith("http"):
                cell.hyperlink = v
                cell.font = Font(name=F, size=10, color="0563C1", underline="single")
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(cols))}{start + len(rows)}"


def okuma_notu(ws):
    lines = [
        ("İstanbul ilçe belediyeleri: asbest verisi taraması (ikinci tur, 13 Ekim 2026)", True),
        ("Birinci tur bulut sunucusundan (26 site), ikinci tur Türkiye IP'li tarayıcıdan (buluttan açılmayan 13 site + rapor sayfası bulunamayan 5 ilçe) yapıldı.", False),
        ("", False),
        ("Bu kitapta ne var", True),
        ("Kapsam 39 ilçe: her belediye için neye bakıldı, ne bulundu, ne kaldı.", False),
        ("Veri noktaları: belgede gözle okunmuş her sayı; kaynak URL, PDF sayfası ve kısa alıntı ile.", False),
        ("İlçe-yıl: veri noktalarının ilçe ve yıla göre düzenlenmiş hali; tespit oranı ve bina başına miktar formülle hesaplanır.", False),
        ("Maltepe mahalle: faaliyet raporlarındaki mahalle bazlı asbest denetim sayıları (2016, 2019).", False),
        ("Kadıköy mahalle tahmini: ilçe asbest serisinin riskli yapı tespitlerine göre mahallelere dağıtılması (MODEL, veri değil).", False),
        ("Bayrampaşa 2025 aylık: yıkım öncesi numune alınan adreslerin aylara dağılımı.", False),
        ("Belgeler (kanıt): indirilen her belgenin erişim durumu, SHA-256 özeti ve erişim zamanı.", False),
        ("", False),
        ("Okurken dikkat", True),
        ("Kaynak türü 'tarayıcıda okundu' olan satırların (Kartal, Şişli, Bayrampaşa) PDF kopyası bu klasörde yok; URL ve sayfa no ile açıp doğrulayın, yayından önce kopyasını alın.", False),
        ("Kartal 2024 raporunda '2023 yılında yıkılan 453 yapı' yazıyor; muhtemelen kopyala-yapıştır hatası, belediyeye sorulmalı.", False),
        ("Göstergeler belediyeler arasında aynı şeyi ölçmüyor: 'denetim', 'envanter raporu', 'araştırma yapılan adres', 'inşaat sahası' farklı sayım birimleri. İlçeleri doğrudan karşılaştırmadan önce tanımı kontrol edin.", False),
        ("Kadıköy'de payda 'denetlenen inşaat sahası', pay 'asbest tespit edilen bina'dır; oran yaklaşık bir göstergedir.", False),
        ("Kağıthane'nin 2020 ve 2021 tonajı kendi raporları arasında çelişkili (34 ve 50 ton; 2024 raporundaki grafikte 11 ve 11 ton). Belediyeye sorulmalı.", False),
        ("Yayına girecek her sayıyı belgenin ilgili sayfasında gözle doğrulayın (Veri noktaları sayfasında URL ve sayfa no var).", False),
        ("Mavi yazılı hücreler belgeden alınan girdilerdir; siyah hücreler formüldür.", False),
    ]
    ws.column_dimensions["A"].width = 140
    for i, (t, b) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(name=F, size=12 if i == 1 else 10, bold=b)
        c.alignment = Alignment(wrap_text=True)


def ilce_yil(ws, pts):
    """İlçe-yıl tablosu: denetim/tespit/miktar girdileri + formüllü oranlar."""
    den_keys = ("denetlenen", "denetimi", "kontrolü yapılan", "araştırması yapılan", "envanter raporu",
                "yönetimi yapılan", "incelenen", "numune alınan")
    tes_keys = ("tespit edilen", "sökümü yaptırılan")
    mik_keys = ("bertaraf", "sökülen", "uzaklaştırılan", "toplanan")
    grid = defaultdict(dict)
    for p in pts:
        if p["duzey"] != "ilçe":
            continue
        g = p["gosterge"]
        k = (p["ilce"], p["yil"])
        if any(x in g for x in tes_keys):
            grid[k].setdefault("tespit", (p["deger"], g, p))
        elif any(x in g for x in mik_keys):
            grid[k].setdefault("miktar", (p["deger"], g, p))
        elif any(x in g for x in den_keys):
            grid[k].setdefault("denetim", (p["deger"], g, p))
    cols = ["İlçe", "Yıl", "Denetlenen / incelenen", "Payda tanımı", "Asbest tespit edilen bina",
            "Tespit oranı", "Miktar", "Birim", "Miktar (kg)", "Tespit edilen bina başına kg", "Kaynak (ilk gösterge)"]
    header(ws, 1, cols, [14, 7, 14, 38, 14, 11, 11, 7, 12, 14, 60])
    for i, k in enumerate(sorted(grid), 2):
        g = grid[k]
        ws.cell(row=i, column=1, value=k[0]).font = BODY
        ws.cell(row=i, column=2, value=str(k[1])).font = BODY
        if "denetim" in g:
            c = ws.cell(row=i, column=3, value=float(g["denetim"][0]) if "." in g["denetim"][0] else int(g["denetim"][0]))
            c.font = INPUT
            ws.cell(row=i, column=4, value=g["denetim"][1]).font = BODY
        if "tespit" in g:
            c = ws.cell(row=i, column=5, value=int(float(g["tespit"][0])))
            c.font = INPUT
        ws.cell(row=i, column=6, value=f'=IF(AND(ISNUMBER(C{i}),ISNUMBER(E{i}),C{i}>0),E{i}/C{i},"")').font = BODY
        ws.cell(row=i, column=6).number_format = "0.0%"
        if "miktar" in g:
            v = g["miktar"][0]
            c = ws.cell(row=i, column=7, value=float(v) if "." in v else int(v))
            c.font = INPUT
            ws.cell(row=i, column=8, value=g["miktar"][2]["birim"]).font = INPUT
        ws.cell(row=i, column=9, value=f'=IF(ISNUMBER(G{i}),IF(H{i}="ton",G{i}*1000,G{i}),"")').font = BODY
        ws.cell(row=i, column=9).number_format = "#,##0"
        ws.cell(row=i, column=10, value=f'=IF(AND(ISNUMBER(I{i}),ISNUMBER(E{i}),E{i}>0),I{i}/E{i},"")').font = BODY
        ws.cell(row=i, column=10).number_format = "#,##0"
        src = next(iter(g.values()))[2]
        c = ws.cell(row=i, column=11, value=src["kaynak_url"])
        if src["kaynak_url"].startswith("http"):
            c.hyperlink = src["kaynak_url"]
            c.font = Font(name=F, size=10, color="0563C1", underline="single")
    ws.auto_filter.ref = f"A1:K{len(grid) + 1}"


def maltepe_mahalle(ws):
    cols = ["Mahalle", "2016 denetim (Mart-Aralık)", "2016 tahmini asbestli bina", "2019 denetim",
            "2019 tahmini asbestli bina"]
    header(ws, 1, cols, [18, 16, 16, 12, 16])
    names = list(MALTEPE_MAH_2016)
    n = len(names)
    for i, m in enumerate(names, 2):
        ws.cell(row=i, column=1, value=m).font = BODY
        ws.cell(row=i, column=2, value=MALTEPE_MAH_2016[m]).font = INPUT
        v19 = MALTEPE_MAH_2019[m]
        if v19 is not None:
            ws.cell(row=i, column=4, value=v19).font = INPUT
        ws.cell(row=i, column=3, value=f"=B{i}*$B${n + 6}").font = BODY
        ws.cell(row=i, column=5, value=f'=IF(ISNUMBER(D{i}),D{i}*$D${n + 6},"")').font = BODY
        for col in (3, 5):
            ws.cell(row=i, column=col).number_format = "0.0"
    t = n + 2
    ws.cell(row=t, column=1, value="Toplam").font = BOLD
    for col in (2, 3, 4, 5):
        L = get_column_letter(col)
        ws.cell(row=t, column=col, value=f"=SUM({L}2:{L}{n + 1})").font = BOLD
    ws.cell(row=t + 2, column=1, value="İlçe: asbest bulunan bina").font = BODY
    ws.cell(row=t + 2, column=2, value=120).font = INPUT
    ws.cell(row=t + 2, column=4, value=20).font = INPUT
    ws.cell(row=t + 3, column=1, value="İlçe: denetlenen bina").font = BODY
    ws.cell(row=t + 3, column=2, value=514).font = INPUT
    ws.cell(row=t + 3, column=4, value=195).font = INPUT
    ws.cell(row=t + 4, column=1, value="İlçe tespit oranı").font = BOLD
    ws.cell(row=t + 4, column=2, value=f"=B{t + 2}/B{t + 3}").number_format = "0.0%"
    ws.cell(row=t + 4, column=4, value=f"=D{t + 2}/D{t + 3}").number_format = "0.0%"
    ws.cell(row=t + 6, column=1, value=("Kaynak: Maltepe Belediyesi 2016 (s.105) ve 2019 (s.165-166) faaliyet raporları. "
                                        "Tahmini sütunlar ilçe tespit oranının her mahalleye eşit uygulanmasıdır (model). "
                                        "2019 tablosunda Gülensu, Büyükbakkalköy ve Başıbüyük için değer yok.")).font = BODY


BAYRAMPASA_2025_AYLIK = [("Şubat", 10), ("Mart", 15), ("Nisan", 8), ("Mayıs", 28), ("Haziran", 11), ("Temmuz", 20),
                         ("Ağustos", 24), ("Eylül", 24), ("Ekim", 44), ("Kasım", 40), ("Aralık", 53)]


def bayrampasa_aylik(ws):
    header(ws, 1, ["Ay (2025)", "Yıkım öncesi numune alınan adres"], [14, 30])
    for i, (ay, n) in enumerate(BAYRAMPASA_2025_AYLIK, 2):
        ws.cell(row=i, column=1, value=ay).font = BODY
        ws.cell(row=i, column=2, value=n).font = INPUT
    t = len(BAYRAMPASA_2025_AYLIK) + 2
    ws.cell(row=t, column=1, value="Toplam").font = BOLD
    ws.cell(row=t, column=2, value=f"=SUM(B2:B{t - 1})").font = BOLD
    ws.cell(row=t + 1, column=1, value="Raporda toplam").font = BODY
    ws.cell(row=t + 1, column=2, value=277).font = INPUT
    ws.cell(row=t + 2, column=1, value="Asbest tespit edilen bina").font = BODY
    ws.cell(row=t + 2, column=2, value=49).font = INPUT
    ws.cell(row=t + 3, column=1, value="Tespit oranı").font = BOLD
    ws.cell(row=t + 3, column=2, value=f"=B{t + 2}/B{t + 1}").number_format = "0.0%"
    ws.cell(row=t + 4, column=1, value="Bertaraf (kg)").font = BODY
    ws.cell(row=t + 4, column=2, value=40500).font = INPUT
    ws.cell(row=t + 5, column=1, value="Tespit edilen bina başına kg").font = BOLD
    ws.cell(row=t + 5, column=2, value=f"=B{t + 4}/B{t + 2}").number_format = "#,##0"
    ws.cell(row=t + 7, column=1, value=(
        "Kaynak: Bayrampaşa Belediyesi 2025 Faaliyet Raporu, PDF s.64 (basılı s.124). Uygulama 14.02.2025'te başladığı "
        "için Ocak yok. Aylık dağılım yıkım ruhsatı başvurusu akışını gösterir; tespit sayısı aylara bölünmemiş.")).font = BODY
    ws.cell(row=t + 8, column=1, value="https://www.bayrampasa.bel.tr/isDosyalar/2026/04/21/is_4692.pdf").font = BODY
    ws.column_dimensions["A"].width = 28


def kadikoy_tahmin(ws):
    rk = [r for r in csv.DictReader(open(RAW / "kadikoy" / "riskli_yapi_mahalle_2012_2022.csv",
                                          encoding="utf-8-sig"), delimiter=";") if r["Mahalle"] != "TOPLAM"]
    years = list(range(2015, 2023))
    cols = ["Mahalle"] + [f"Riskli yapı {y}" for y in years] + ["Toplam yapı", "Tahmini asbestli bina 2015-22",
                                                                   "Tahmini ton 2015-22", "Bin yapıda tahmini asbestli"]
    header(ws, 1, cols, [16] + [10] * len(years) + [10, 16, 14, 14])
    n = len(rk)
    tot_r, det_r, ton_r = n + 3, n + 4, n + 5
    first, last = get_column_letter(2), get_column_letter(1 + len(years))
    for i, r in enumerate(rk, 2):
        ws.cell(row=i, column=1, value=r["Mahalle"]).font = BODY
        for j, y in enumerate(years, 2):
            ws.cell(row=i, column=j, value=int(r[str(y)])).font = INPUT
        c0 = 2 + len(years)
        ws.cell(row=i, column=c0, value=int(r["Toplam_Yapi"])).font = INPUT
        rng = f"{first}{i}:{last}{i}"
        ws.cell(row=i, column=c0 + 1,
                value=f"=SUMPRODUCT({rng}/{first}${tot_r}:{last}${tot_r}*{first}${det_r}:{last}${det_r})").number_format = "0.0"
        ws.cell(row=i, column=c0 + 2,
                value=f"=SUMPRODUCT({rng}/{first}${tot_r}:{last}${tot_r}*{first}${ton_r}:{last}${ton_r})").number_format = "0.0"
        ws.cell(row=i, column=c0 + 3,
                value=f"=1000*{get_column_letter(c0 + 1)}{i}/{get_column_letter(c0)}{i}").number_format = "0.0"
    ws.cell(row=tot_r, column=1, value="Toplam riskli yapı").font = BOLD
    ws.cell(row=det_r, column=1, value="İlçe: asbest tespit edilen bina").font = BOLD
    ws.cell(row=ton_r, column=1, value="İlçe: bertaraf ton").font = BOLD
    for j, y in enumerate(years, 2):
        L = get_column_letter(j)
        ws.cell(row=tot_r, column=j, value=f"=SUM({L}2:{L}{n + 1})").font = BOLD
        ws.cell(row=det_r, column=j, value=KADIKOY[y][1]).font = INPUT
        ws.cell(row=ton_r, column=j, value=KADIKOY[y][2]).font = INPUT
    ws.cell(row=ton_r + 2, column=1, value=(
        "Kaynak: Kadıköy Açık Veri Portalı — 'Asbest Verileri' ve 'Riskli Yapı Tespiti Yapılan Binaların Mahalle "
        "Dağılımları 2012-2022'. MODEL: her yılın ilçe tespiti, o yılın riskli yapı tespitine orantılı dağıtıldı; "
        "tespit olasılığı mahalleler arasında eşit varsayıldı. Riskli yapı tespiti ile yıkım arasındaki gecikme ve "
        "bina yaşı farkları sonucu değiştirir.")).font = BODY
    ws.cell(row=ton_r + 3, column=1, value=KADIKOY_CSV).font = BODY


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Okuma notu"
    okuma_notu(ws)
    pts = load("veri_noktalari.csv")
    table(wb.create_sheet("Kapsam 39 ilçe"), load("kapsam_39.csv"),
          ["ilce", "yaka", "erisim", "site", "site_haritasi_url", "site_ici_arama", "acik_veri",
           "taranan_faaliyet_raporu_yillari", "indirilen_belge", "asbest_gecen_belge", "ocr_gereken", "sonuc",
           "not", "dizinde_gorulen", "sonraki_adim"],
          [13, 8, 30, 28, 10, 16, 34, 28, 9, 9, 8, 36, 50, 40, 50], link_col="site")
    table(wb.create_sheet("Veri noktaları"), pts,
          ["ilce", "duzey", "mahalle", "yil", "gosterge", "deger", "birim", "kaynak_turu", "kaynak_url", "sayfa",
           "alinti", "not"], [12, 8, 14, 6, 34, 10, 7, 14, 50, 6, 70, 30], link_col="kaynak_url")
    ilce_yil(wb.create_sheet("İlçe-yıl"), pts)
    maltepe_mahalle(wb.create_sheet("Maltepe mahalle"))
    kadikoy_tahmin(wb.create_sheet("Kadıköy mahalle tahmini"))
    bayrampasa_aylik(wb.create_sheet("Bayrampaşa 2025 aylık"))
    table(wb.create_sheet("Belgeler (kanıt)"), load("belgeler.csv"),
          ["ilce", "url", "baslik", "durum", "http", "tur", "bytes", "sayfa_sayisi", "asbest_gecisi", "ocr_gerekli",
           "sha256", "erisim"], [12, 60, 30, 12, 6, 6, 11, 7, 7, 7, 20, 20], link_col="url")
    wb["İlçe-yıl"]["F1"].comment = Comment("Tespit edilen bina / denetlenen (veya incelenen) birim. "
                                           "Kadıköy'de payda inşaat sahasıdır.", "tarama")
    for sh in wb.worksheets:  # tüm hücrelerde Arial
        for row in sh.iter_rows():
            for c in row:
                if c.font is None or c.font.name != F:
                    c.font = Font(name=F, size=c.font.size or 10, bold=c.font.bold, italic=c.font.italic,
                                  color=c.font.color, underline=c.font.underline)
    out = XLSX / "asbest_istanbul_ikinci_tur.xlsx"
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
