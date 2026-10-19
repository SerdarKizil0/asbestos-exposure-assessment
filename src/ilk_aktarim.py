"""TEK SEFERLİK AKTARIM (20 Ekim 2026): eski Excel kitaplarını ve dağınık kayıtları CSV'ye çevirir.

Girdiler (git'e girmez):
  gelen/excel/asbest_yayilim_veri_tablosu.xlsx      literatür kaynakları, parametreler, ölçümler, model kurulumları, hesaplar
  gelen/excel/asbest_yayilim_standart_degerler.xlsx  modelin standart değer seti ve açıklamaları
  gelen/kopyalar/kadikoy_asbest_verileri.xlsx        Kadıköy belgelerinin URL, erişim zamanı ve SHA-256 kaydı
  ciktilar/tarama/belgeler.csv, veri/dogrulanmis/veri_noktalari.csv, veri/ham/**/*.json, veri/ham/log/istekler.jsonl

Çıktılar:
  kaynaklar.csv                       her belgenin URL, erişim tarihi, SHA-256, lisans ve arşiv bağlantısı
  veri/parametreler.csv               literatür parametreleri (P01–P38)
  veri/literatur/*.csv                ölçümler, model kurulumları, model sonuçları, hesaplar, standart değerler, kaynak rolleri

Bu betik yeniden çalıştırılırsa kaynaklar.csv'deki elle yapılmış düzeltmeleri ezer. Aktarımdan sonra
kaynaklar.csv elle (ya da src/kaynak_ekle.py ile) güncellenir; Excel kitapları src/excel_uret.py ile üretilir.

Kullanım:  python src/ilk_aktarim.py
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import zipfile
from pathlib import Path

import openpyxl

KOK = Path(__file__).resolve().parents[1]
GELEN = KOK / "gelen"
LIT = KOK / "veri" / "literatur"

# Eski kitaptaki "Doğrulama durumu" → notu yazanın kaynağı ne kadar okuduğu (bizim doğrulamamız değildir).
OKUMA = {
    "Tam metin okundu": "tam_metin",
    "Tam metin okundu (ikincil derleme)": "ikincil",
    "Özet okundu": "ozet",
    "Özet bildirinin tamamı okundu": "tam_metin",
    "Kurum basın özeti okundu": "ikincil",
    "Özet, ikincil sayfada okundu": "ikincil",
    "İkincil özet okundu": "ikincil",
    "Sayfa açıldı, metin görüntülenemedi": "acilamadi",
    "Açılamadı (yalnızca yayın kaydı özeti görüldü)": "acilamadi",
    "Açılamadı (ikincil aktarım)": "acilamadi",
}

# Eski kitaptaki Kaynak No → kaynak_id
LIT_ID = {
    1: "terazono2000", 2: "gilliam2005", 3: "pession2011", 4: "stenchikov2006", 5: "atsdr2003",
    6: "moorcroft1984", 7: "lioy2002", 8: "springston2011", 9: "swayze2006", 10: "yamamoto1999",
    11: "epa2001", 12: "inui2012", 13: "baran2024", 14: "oral2025", 15: "beck2003", 16: "stefani2005",
    17: "neitzel2020", 18: "dukeforest1996", 19: "ravina2022", 20: "hochrainer1989", 21: "kottek_yuen2022",
    22: "lorber2007", 23: "nrc2003", 24: "ervik2021", 25: "ers2026", 26: "tema2024",
}

IBB_LISANS = ("İBB Açık Veri Lisansı (kopyalama, uyarlama ve ticari kullanım serbest; atıf zorunlu; "
              "CC BY 4.0 ile uyumlu: data.ibb.gov.tr/license, 20.10.2026 okundu)")

KAYNAK_ALANLAR = ["kaynak_id", "tur", "kunye", "yayinci", "yil", "url", "erisim_tarihi", "sha256", "lisans",
                  "dagitim", "arsiv_url", "okuma", "yerel_ad", "not"]


def yaz(yol: Path, satirlar: list[dict], alanlar: list[str]) -> None:
    yol.parent.mkdir(parents=True, exist_ok=True)
    with open(yol, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=alanlar, lineterminator="\n")
        w.writeheader()
        for s in satirlar:
            w.writerow({k: hucre(s.get(k)) for k in alanlar})
    print(f"  {yol.relative_to(KOK)}: {len(satirlar)} satır")


def hucre(v):
    """CSV'ye yazılacak biçim: ondalık nokta, boş = belgede yok."""
    if v is None:
        return ""
    if isinstance(v, float):
        return repr(v)
    return v


def tablo(ws, baslik_satiri: int = 1, son: int | None = None) -> tuple[list[str], list[list]]:
    rows = list(ws.iter_rows(min_row=baslik_satiri, max_row=son or ws.max_row, values_only=True))
    basliklar = [h for h in rows[0] if h is not None]
    veri = [list(r[:len(basliklar)]) for r in rows[1:] if any(v is not None for v in r)]
    return basliklar, veri


def sha256(yol: Path) -> str:
    h = hashlib.sha256()
    with open(yol, "rb") as f:
        for parca in iter(lambda: f.read(1 << 20), b""):
            h.update(parca)
    return h.hexdigest()


# ---------------------------------------------------------------- literatür kitabı
def literatur() -> list[dict]:
    wb = openpyxl.load_workbook(GELEN / "excel" / "asbest_yayilim_veri_tablosu.xlsx")
    print("asbest_yayilim_veri_tablosu.xlsx")

    # Kaynaklar → kaynak rolleri (künye ve bağlantı kaynaklar.csv'de tutulur)
    _, kay = tablo(wb["Kaynaklar"])
    roller, kaynaklar = [], []
    for no, kisa, kunye, olay, rol, diger, sayfalar, durum, url in kay:
        kid = LIT_ID[no]
        roller.append({"no": no, "kaynak_id": kid, "kisa_ad": kisa, "olay_konu": olay, "birincil_rol": rol,
                       "diger_roller": diger, "excel_sayfalari": sayfalar, "okuma_durumu": durum})
        yil = re.search(r"\((\d{4})\)", kunye)
        kaynaklar.append({"kaynak_id": kid, "tur": "literatur", "kunye": kunye, "yil": yil.group(1) if yil else "",
                          "url": url, "lisans": "yayıncı lisansı; dağıtılmaz", "dagitim": "yalniz_baglanti",
                          "okuma": OKUMA[durum],
                          "not": f"{durum}. Tablo 5 Ekim 2026'da derlendi; erişim tarihi kaydedilmemiş. "
                                 "Birincil belge depoda yok, değerler belgeyle karşılaştırılmadı."})
    yaz(LIT / "kaynak_rolleri.csv", roller,
        ["no", "kaynak_id", "kisa_ad", "olay_konu", "birincil_rol", "diger_roller", "excel_sayfalari", "okuma_durumu"])
    okuma = {r["kaynak_id"]: OKUMA[r["okuma_durumu"]] for r in roller}

    # Parametreler
    _, par = tablo(wb["Parametreler"])
    satirlar = []
    for pid, grup, ad, deger, alt, ust, birim, kosul, kno, kaynak, yer, nott in par:
        satirlar.append({"parametre_id": pid, "grup": grup, "parametre": ad, "deger": deger, "alt": alt, "ust": ust,
                         "birim": birim, "kosul": kosul, "kaynak_no": kno, "kaynak_id": LIT_ID[kno], "kaynak": kaynak,
                         "yer": yer, "okuma": okuma[LIT_ID[kno]], "dogrulandi": "hayir", "not": nott})
    yaz(KOK / "veri" / "parametreler.csv", satirlar,
        ["parametre_id", "grup", "parametre", "deger", "alt", "ust", "birim", "kosul", "kaynak_no", "kaynak_id",
         "kaynak", "yer", "okuma", "dogrulandi", "not"])

    # Ölçümler
    b, olc = tablo(wb["Ölçümler"])
    satirlar = []
    for r in olc:
        d = dict(zip(b, r))
        satirlar.append({"olcum_id": d["ID"], "olay": d["Olay"], "yer_nokta": d["Yer / nokta"], "donem": d["Dönem"],
                         "olculen": d["Ölçülen"], "yontem": d["Yöntem"], "deger": d["Değer"], "alt": d["Alt"],
                         "ust": d["Üst"], "birim": d["Birim"], "n": d["N (örnek/nokta)"], "kaynak_no": d["Kaynak No"],
                         "kaynak_id": LIT_ID[d["Kaynak No"]], "kaynak": d["Kaynak"], "yer": d["Yer"],
                         "okuma": okuma[LIT_ID[d["Kaynak No"]]], "dogrulandi": "hayir", "not": d["Not"]})
    yaz(LIT / "olcumler.csv", satirlar, list(satirlar[0].keys()))

    # Model kurulumları
    b, mk = tablo(wb["Model_Kurulumları"])
    alan = ["calisma", "olay", "model", "alan_olcek", "yatay_cozunurluk", "zaman_cozunurlugu", "kaynak_tanimi", "salim",
            "meteoroloji_girdisi", "cokelme", "gecerlilik_kisit", "dogrulama_verisi", "yer", "kaynak_no"]
    satirlar = [dict(zip(alan, r)) | {"kaynak_id": LIT_ID[r[-1]]} for r in mk]
    yaz(LIT / "model_kurulumlari.csv", satirlar, alan + ["kaynak_id"])

    # Model sonuçları: iki ayrı tablo (A: WTC seyrelme, B: Kobe karşılaştırması)
    ws = wb["Model_Sonuçları"]
    alan = ["tarih_donem", "meteorolojik_durum", "karisim_alt_m", "karisim_ust_m", "ruzgar_alt_ms", "ruzgar_ust_ms",
            "seyrelme_alt", "seyrelme_ust", "mesafe_alt_km", "mesafe_ust_km", "gozlem_pm25", "not"]
    a = [dict(zip(alan, r)) | {"kaynak_id": "gilliam2005"}
         for r in ws.iter_rows(min_row=3, max_row=18, max_col=12, values_only=True)]
    yaz(LIT / "wtc_seyrelme.csv", a, alan + ["kaynak_id"])
    alan = ["kalem", "aciklama", "deger", "alt", "ust", "birim"]
    bb = [dict(zip(alan, r)) | {"kaynak_id": "terazono2000"}
          for r in ws.iter_rows(min_row=23, max_row=28, max_col=6, values_only=True)]
    yaz(LIT / "kobe_karsilastirma.csv", bb, alan + ["kaynak_id"])

    # Hesaplar: hücre düzeni korunur (formüller satır numarasıyla birbirine bağlı)
    ws = wb["Hesaplar"]
    satirlar = []
    for r in range(1, ws.max_row + 1):
        hucreler = [ws.cell(r, c) for c in range(1, 7)]
        if all(h.value is None for h in hucreler):
            continue
        a_ = hucreler[0]
        dolgu = a_.fill.fgColor.rgb if a_.fill.fill_type else None
        if r == 1:
            tur = "baslik"
        elif dolgu == "FFBDD7EE":
            tur = "bolum"
        elif dolgu == "FFD9D9D9":
            tur = "sutun_basligi"
        else:
            tur = "kalem"
        b_ = hucreler[1]
        if tur != "kalem" or b_.value is None:
            girdi = ""
        elif isinstance(b_.value, str) and b_.value.startswith("="):
            girdi = "formul"
        elif b_.fill.fill_type and b_.fill.fgColor.rgb == "FFFFFF00":
            girdi = "varsayim"
        else:
            girdi = "kaynak"
        satirlar.append({"satir": r, "tur": tur, "kalem": a_.value, "deger": b_.value, "birim": hucreler[2].value,
                         "makale_degeri": hucreler[3].value, "fark": hucreler[4].value, "not": hucreler[5].value,
                         "girdi": girdi, "deger_bicimi": b_.number_format, "makale_bicimi": hucreler[3].number_format})
    yaz(LIT / "hesaplar.csv", satirlar, list(satirlar[0].keys()))
    return kaynaklar


# ---------------------------------------------------------------- standart değerler
DAYANAK = {  # standart değer satırı → dayandığı literatür satırı
    "Binadaki asbest yükü (kat alanı başına)": "P01",
    "Havaya karışan pay (salım faktörü)": "P20",
    "Kütleden lif sayısına çevirme": "P26",
    "Lif çökelme hızı, 0,5 µm lif (pratikte ihmal edilir)": "hesaplar.csv satır 61",
    "Karışım yüksekliği – durgun gece (en kötü durum)": "wtc_seyrelme.csv",
    "Karışım yüksekliği – gündüz": "wtc_seyrelme.csv",
    "Rüzgâr hızı – durgun gece (en kötü durum)": "wtc_seyrelme.csv",
    "Model düzeltme katsayısı (kaynağa 200 m'den yakın)": "P34, P35",
    "Olağan şehir havası": "M03, M04",
    "Deprem sonrası kent havası (Kobe)": "M01, M02",
    "Asbesti sökülmeden yıkılan binanın sınırı (Kobe)": "M17, M18",
    "Türkiye işyeri sınır değeri (8 saatlik ortalama)": "asbest_yonetmeligi2013",
    "Asbestin ölçüldüğü en uzak mesafe (11 Eylül)": "M32",
}
STD_ID = {"Terazono ve ark. 2000": "terazono2000", "Gilliam ve ark. 2005": "gilliam2005",
          "Moorcroft ve Duggan 1984": "moorcroft1984", "Duke Forest deneyi 1996": "dukeforest1996",
          "Swayze ve ark. 2006": "swayze2006", "Asbest Yönetmeliği 2013": "asbest_yonetmeligi2013"}


def sayisal(gosterim: str) -> tuple:
    """'0,086' → (0.086, None, None); '100–200' → (None, 100, 200); '%0,01' → (0.0001, ...)."""
    s = gosterim.replace(".", "").replace(",", ".").strip()
    yuzde = s.startswith("%")
    s = s.lstrip("%").lstrip(">").strip()
    s = re.sub(r"\s*\(.*\)$", "", s)
    m = re.fullmatch(r"([\d.]+)\s*–\s*([\d.]+)", s)
    if m:
        return None, tam(float(m.group(1))), tam(float(m.group(2)))
    v = float(s)
    v = v / 100 if yuzde else v
    if gosterim.strip().startswith(">"):
        return None, tam(v), None
    return tam(v), None, None


def tam(v: float):
    return int(v) if float(v).is_integer() else v


def standart_degerler() -> None:
    print("asbest_yayilim_standart_degerler.xlsx")
    wb = openpyxl.load_workbook(GELEN / "excel" / "asbest_yayilim_standart_degerler.xlsx")
    ws = wb["Standart Değerler"]
    aciklama = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(23, 35)}
    satirlar, bolum = [], None
    for r in range(5, 21):
        a, b, c, d = (ws.cell(r, k).value for k in range(1, 5))
        if a and b is None:
            bolum = a
            continue
        bas = next(k for k in aciklama if a.startswith(k) or k.startswith(a.split(" (")[0].split(",")[0]))
        deger, alt, ust = sayisal(b)
        kisa = re.sub(r"\s*\(.*\)$", "", d)
        satirlar.append({"bolum": bolum, "parametre": a, "gosterim": b, "deger": deger, "alt": alt, "ust": ust,
                         "birim": c, "kaynak": d, "kaynak_id": STD_ID.get(kisa, STD_ID.get(d.split(",")[0], "")),
                         "dayanak": DAYANAK[a], "aciklama_basligi": bas, "aciklama": aciklama[bas]})
    yaz(LIT / "standart_degerler.csv", satirlar, list(satirlar[0].keys()))

    # Kitabın kendi Kaynaklar sayfası: künye kaynaklar.csv'den gelir, yalnız bu kitaba özgü ek not burada tutulur.
    ws = wb["Kaynaklar"]
    temel = {k["kaynak_id"]: k["kunye"] for k in literatur_kunyeleri() + raporlar()}
    satirlar = []
    for kisa, kunye in ws.iter_rows(min_row=2, max_col=2, values_only=True):
        kid = STD_ID[re.sub(r"\s*\(USGS\)$", "", kisa)]
        baz = temel[kid]
        if not kunye.startswith(baz):
            raise SystemExit(f"Künye kaynaklar.csv ile uyuşmuyor: {kisa}")
        satirlar.append({"kisa_ad": kisa, "kaynak_id": kid, "ek_not": kunye[len(baz):].strip()})
    yaz(LIT / "standart_degerler_kaynaklari.csv", satirlar, ["kisa_ad", "kaynak_id", "ek_not"])


def literatur_kunyeleri() -> list[dict]:
    ws = openpyxl.load_workbook(GELEN / "excel" / "asbest_yayilim_veri_tablosu.xlsx")["Kaynaklar"]
    return [{"kaynak_id": LIT_ID[r[0]], "kunye": r[2]} for r in ws.iter_rows(min_row=2, max_col=3, values_only=True) if r[0]]


# ---------------------------------------------------------------- kaynaklar.csv
def istek_zamanlari() -> dict:
    """İstek günlüğünden her URL'nin ilk başarılı erişim zamanı."""
    z = {}
    yol = KOK / "veri" / "ham" / "log" / "istekler.jsonl"
    for satir in open(yol, encoding="utf-8"):
        r = json.loads(satir)
        if r.get("status") == 200 or r.get("durum") == "ok":
            z.setdefault(r["url"], r["ts"])
    return z


def belediye_belgeleri(zaman: dict) -> list[dict]:
    """Veri noktalarının dayandığı belediye belgeleri (tarama kaydıyla birleştirilmiş)."""
    nokta = list(csv.DictReader(open(KOK / "veri" / "dogrulanmis" / "veri_noktalari.csv", encoding="utf-8-sig")))
    belge = {d["url"]: d for d in csv.DictReader(open(KOK / "ciktilar" / "tarama" / "belgeler.csv", encoding="utf-8-sig"))}
    slug = {s["ilce"]: s["slug"] for s in csv.DictReader(open(KOK / "src" / "tarama" / "seeds" / "belediyeler.csv", encoding="utf-8"))}
    goruldu, cikti = set(), []
    for p in nokta:
        url = p["kaynak_url"]
        if url in goruldu or "acikveri.kadikoy" in url:
            continue
        goruldu.add(url)
        d = belge.get(url, {})
        # Rapor yılı: önce tarama kaydındaki başlık ("Faaliyet raporu 2017 | …"), yoksa veri noktasının yılı.
        # URL'deki yıl yanıltıcı olabilir (yükleme tarihi ya da "2010-2019" gibi klasör adı).
        m = re.search(r"Faaliyet raporu (\d{4})", d.get("baslik", ""))
        yil = m.group(1) if m else p["yil"][:4]
        haber = "haber" in p["kaynak_turu"]
        kid = f"{slug[p['ilce']]}_{'haber' if haber else 'fr' + yil}"
        if haber or any(k["kaynak_id"] == kid for k in cikti):
            kid += "_" + hashlib.sha1(url.encode()).hexdigest()[:6]
        kopya = bool(d.get("sha256"))
        cikti.append({
            "kaynak_id": kid, "tur": "belediye_haberi" if haber else "faaliyet_raporu",
            "kunye": f"{p['ilce']} Belediyesi, {d.get('baslik') or p['kaynak_turu']}", "yayinci": f"{p['ilce']} Belediyesi",
            "yil": yil, "url": url,
            "erisim_tarihi": d.get("erisim") or zaman.get(url) or "2026-10-13 (tarayıcı)",
            "sha256": d.get("sha256", ""), "lisans": "belediye yayını; dağıtılmaz", "dagitim": "yalniz_baglanti",
            "okuma": "tam_metin", "yerel_ad": f"veri/ham/{slug[p['ilce']]}/" if kopya else "",
            "not": "Sayılar src/tarama/veri_noktalari.py ile belgede bulunup sayfa ve alıntıyla bağlandı."
                   if kopya else "Türkiye IP'li tarayıcıda okundu; yerel kopya ve SHA-256 yok (harvest.py ile indirilmeli).",
        })
    return cikti


def kadikoy_belgeleri() -> list[dict]:
    ws = openpyxl.load_workbook(GELEN / "kopyalar" / "kadikoy_asbest_verileri.xlsx")["Kaynaklar"]
    ad = {"Kadıköy Açık Veri: Asbest Verileri (CSV)": "kadikoy_acikveri_asbest", "2018 Faaliyet Raporu": "kadikoy_fr2018",
          "2021 Faaliyet Raporu": "kadikoy_fr2021", "2022 Faaliyet Raporu": "kadikoy_fr2022",
          "2023 Faaliyet Raporu": "kadikoy_fr2023", "2024 Faaliyet Raporu": "kadikoy_fr2024"}
    cikti = []
    for belge, sayfa, url, erisim, sha, nott in ws.iter_rows(min_row=2, values_only=True):
        kid = ad.get(belge) or "kadikoy_" + re.sub(r"[^a-z0-9]+", "_", belge.lower().translate(
            str.maketrans("çğıöşüâ", "cgiosua")))[:40].strip("_")
        acik = "acikveri" in url
        cikti.append({
            "kaynak_id": kid, "tur": "acik_veri" if acik else ("web_sayfasi" if url.endswith(("-5150", "-5161", "-4885")) or "icerik" in url else "belediye_belgesi"),
            "kunye": f"Kadıköy Belediyesi, {belge}", "yayinci": "Kadıköy Belediyesi", "url": url,
            "erisim_tarihi": str(erisim), "sha256": sha,
            "lisans": "CC BY-NC 3.0 (CKAN kaydı: creativecommons.org/licenses/by-nc/3.0/au)" if acik else "belediye yayını; dağıtılmaz",
            "dagitim": "veri/dogrulanmis içinde türetilmiş satırlar" if acik else "yalniz_baglanti",
            "okuma": "tam_metin", "yerel_ad": "veri/ham/kadikoy/",
            "not": "; ".join(x for x in (f"kullanılan sayfa: {sayfa}" if sayfa != "tümü" else "", nott or "") if x),
        })
    return cikti


def acik_veri(zaman: dict) -> list[dict]:
    ham = KOK / "veri" / "ham"
    cikti = []
    m = json.loads((ham / "ibb" / "deprem_senaryosu_analiz_sonuclari.meta.json").read_text(encoding="utf-8"))
    cikti.append({"kaynak_id": "ibb_deprem_senaryosu", "tur": "acik_veri",
                  "kunye": "İBB Açık Veri Portalı, Deprem Senaryosu Analiz Sonuçları (Mw 7,5 gece senaryosu, 959 mahalle)",
                  "yayinci": "İBB Deprem Risk Yönetimi ve Kentsel İyileştirme Daire Başkanlığı", "yil": "2021",
                  "url": m["url"], "erisim_tarihi": m["erisim"], "sha256": m["sha256"], "lisans": IBB_LISANS,
                  "dagitim": "yalniz_baglanti", "okuma": "tam_metin", "yerel_ad": "veri/ham/ibb/deprem_senaryosu_analiz_sonuclari.csv",
                  "not": "Kaynak çalışma KRDAE 2019 (Tablo 4-2 ve 4-19 ile toplamlar %0,04 içinde; Avcılar ve Fatih kitapçıklarıyla "
                         "mahalle düzeyinde birebir). Veri seti oluşturma 19.03.2021, güncelleme 02.03.2023. Windows-1254, ';' ayırıcı."})
    r = json.loads((ham / "ibb" / "resource.json").read_text(encoding="utf-8"))["result"]
    yol = ham / "ibb" / "mahalle_bina_2017.csv"
    cikti.append({"kaynak_id": "ibb_bina_2017", "tur": "acik_veri",
                  "kunye": "İBB Açık Veri Portalı, 2017 Yılı Mahalle Bazlı Bina Sayıları (yaş ve kat sınıfları)",
                  "yayinci": "İBB Deprem Risk Yönetimi ve Kentsel İyileştirme Daire Başkanlığı", "yil": "2021",
                  "url": r["url"], "erisim_tarihi": zaman.get(r["url"], "kaydedilmemiş"), "sha256": sha256(yol),
                  "lisans": IBB_LISANS, "dagitim": "yalniz_baglanti", "okuma": "tam_metin",
                  "yerel_ad": "veri/ham/ibb/mahalle_bina_2017.csv",
                  "not": "Mahalle adı '19 MAYIS' dosyada '19.May' olarak bozulmuş; UAVT koduyla birleştirin."})
    m = json.loads((ham / "ibb" / "ibb_lokasyon.meta.json").read_text(encoding="utf-8"))
    cikti.append({"kaynak_id": "ibb_lokasyon", "tur": "acik_veri", "kunye": "İBB Açık Veri Portalı, İBB lokasyonları",
                  "yayinci": "İBB", "url": m["url"], "erisim_tarihi": m["erisim"], "sha256": m["sha256"],
                  "lisans": IBB_LISANS, "dagitim": "yalniz_baglanti", "okuma": "tam_metin",
                  "yerel_ad": "veri/ham/ibb/ibb_lokasyon.csv", "not": "Modelde kullanılmıyor."})
    lis = {}
    for p in json.loads((ham / "kadikoy" / "model" / "paketler.json").read_text(encoding="utf-8")):
        if isinstance(p, dict) and "name" in p:
            lis[p["name"]] = f"{p.get('license_title')} ({p.get('license_url')})"
    for d in json.loads((ham / "kadikoy" / "model" / "indirilenler.json").read_text(encoding="utf-8")):
        ad = d["dosya"].rsplit(".", 1)[0]
        cikti.append({"kaynak_id": "kadikoy_" + ad.replace("-", "_")[:45], "tur": "acik_veri",
                      "kunye": f"Kadıköy Açık Veri, {ad}", "yayinci": "Kadıköy Belediyesi", "url": d["url"],
                      "erisim_tarihi": d["erisim"], "sha256": d["sha256"], "lisans": lis.get(ad, "kaydedilmemiş"),
                      "dagitim": "model/girdi içinde türetilmiş satırlar", "okuma": "tam_metin",
                      "yerel_ad": f"veri/ham/kadikoy/model/{d['dosya']}",
                      "not": "0-14 yaş dosyası tutarsız, modelde kullanılmadı." if "0-14" in ad else ""})
    url = ("https://acikveri.kadikoy.bel.tr/dataset/834a160d-3c0b-42b8-baca-f6b16fe75738/resource/"
           "0a8282aa-d1ef-4548-a786-7cb9f1ba109b/download/kadkoy-riskli-yap-tespiti-yaplan-binalarn-mahalle-dalmlar-2012-2022.csv")
    cikti.append({"kaynak_id": "kadikoy_riskli_yapi_2012_2022", "tur": "acik_veri",
                  "kunye": "Kadıköy Açık Veri, Riskli Yapı Tespiti Yapılan Binaların Mahalle Dağılımları 2012-2022",
                  "yayinci": "Kadıköy Belediyesi", "url": url, "erisim_tarihi": zaman.get(url, "kaydedilmemiş"),
                  "sha256": sha256(ham / "kadikoy" / "riskli_yapi_mahalle_2012_2022.csv"), "lisans": "CC BY-NC 3.0 (CKAN kaydı: creativecommons.org/licenses/by-nc/3.0/au; 20.10.2026 package_show ile okundu)",
                  "dagitim": "model/girdi içinde türetilmiş satırlar", "okuma": "tam_metin",
                  "yerel_ad": "veri/ham/kadikoy/riskli_yapi_mahalle_2012_2022.csv", "not": ""})
    m = json.loads((ham / "meteo" / "era5_kadikoy_2015_2024.meta.json").read_text(encoding="utf-8"))
    cikti.append({"kaynak_id": "era5_openmeteo_kadikoy", "tur": "acik_veri",
                  "kunye": "Hersbach ve ark. (2020), ERA5 saatlik yeniden analiz; Open-Meteo Historical Weather API üzerinden "
                           "(40,985 K, 29,06 D; 2015-2024)", "yayinci": "Copernicus C3S / Open-Meteo", "url": m["url"],
                  "erisim_tarihi": m["erisim"], "sha256": m["sha256"], "lisans": "CC BY 4.0", "dagitim": "model/girdi içinde frekans tabloları",
                  "okuma": "tam_metin", "yerel_ad": "veri/ham/meteo/era5_kadikoy_2015_2024.json",
                  "not": "Open-Meteo robots.txt tüm yolları kapatıyor; belgelenmiş API olduğu için yalnız bu istekte "
                         "robots kontrolü kapatıldı (model/indir.py). İstenirse Copernicus CDS'ten indirilebilir."})
    cikti.append({"kaynak_id": "osm_istanbul_geojson", "tur": "acik_veri",
                  "kunye": "OpenStreetMap katkıcıları; sahircansurmeli/istanbul-geojson (mahalle_geojson.json), commit c9bf151",
                  "yayinci": "OpenStreetMap", "url": "https://github.com/sahircansurmeli/istanbul-geojson",
                  "erisim_tarihi": "2026-10-14 (model kurulumu)", "lisans": "ODbL 1.0", "dagitim": "model/girdi içinde türetilmiş satırlar",
                  "okuma": "tam_metin", "yerel_ad": "veri/ham/osm/istanbul-geojson/ (yerelde yok, klonlanmalı)",
                  "not": "Kadıköy poligonlarına ad eşleştirmek ve komşu ilçe mahallelerini bulmak için kullanıldı."})
    return cikti


def raporlar() -> list[dict]:
    """Depoya girmeyen yerel PDF'ler (gelen/ ve arşivdeki kopyalar). SHA-256 arşiv envanterinden."""
    env = {}
    with zipfile.ZipFile(GELEN / "arsivler" / "deprem_asbest_tum_veriler_2026-10-19.zip") as z:
        metin = z.read("ENVANTER.csv").decode("utf-8-sig")
    for r in csv.DictReader(io.StringIO(metin)):
        env[Path(r["arsiv_yolu"]).name] = r["sha256"]
    assert env["IBB_Kandilli_2019_ana_rapor.pdf"] == sha256(GELEN / "belgeler" / "IBB_Kandilli_2019_ana_rapor.pdf")
    ortak = {"dagitim": "yalniz_baglanti", "okuma": "tam_metin"}
    return [
        {"kaynak_id": "krdae2019", "tur": "rapor", "yil": "2019", "kunye":
            "İBB ve Boğaziçi Üniversitesi KRDAE Deprem Mühendisliği ABD (2019). İstanbul İli Olası Deprem Kayıp Tahminlerinin "
            "Güncellenmesi Projesi.", "yayinci": "İBB Deprem ve Zemin İnceleme Müdürlüğü",
         "url": "https://depremzemin.ibb.istanbul/tr/istanbul-ili-olasi-deprem-kayip-tahminlerinin-guncellenmesi-projesi-2019",
         "erisim_tarihi": "kaydedilmemiş (dosya tarihi 2026-09-20)", "sha256": env["IBB_Kandilli_2019_ana_rapor.pdf"],
         "lisans": "telifli; dağıtılmaz", "yerel_ad": "gelen/belgeler/IBB_Kandilli_2019_ana_rapor.pdf",
         "not": "342 sayfa; PDF sayfası = basılı sayfa + 36. Proje sayfasında PDF bağlantısı yok (20.10.2026); rapor listesi: https://depremzemin.ibb.istanbul/tr/deprem-raporlari"} | ortak,
        {"kaynak_id": "oyo2009", "tur": "rapor", "yil": "2009", "kunye":
            "OYO International Corporation (Kasım 2009). İstanbul'un Olası Deprem Kayıpları Tahminlerinin Güncellenmesi İşi Sonuç Raporu.",
         "yayinci": "İBB", "url": "", "erisim_tarihi": "kaydedilmemiş (dosya tarihi 2026-09-20)", "sha256": env["ana.pdf"],
         "lisans": "telifli; dağıtılmaz", "yerel_ad": "gelen/arsivler/deprem_asbest_tum_veriler_2026-10-19.zip:02_kaynak_raporlar/ana.pdf",
         "not": "Arşivde 'ana.pdf' adıyla duruyor; 2019 raporu değildir. URL bulunmalı."} | ortak,
        {"kaynak_id": "ibb_kitapcik_avcilar2020", "tur": "rapor", "yil": "2020", "kunye":
            "İBB DEZİM ve BÜ KRDAE (Haziran 2020). İstanbul İli Avcılar İlçesi Olası Deprem Kayıp Tahminleri Kitapçığı.",
         "yayinci": "İBB", "url": "https://depremzemin.ibb.istanbul/tr/olasi-deprem-kayip-tahminleri-ilce-kitapciklari",
         "erisim_tarihi": "kaydedilmemiş (dosya tarihi 2026-09-20)", "sha256": env["avcilar.pdf"], "lisans": "telifli; dağıtılmaz",
         "yerel_ad": "arşiv: 02_kaynak_raporlar/avcilar.pdf",
         "not": "Tablo 5-1 (s.42) İBB senaryo CSV'siyle mahalle düzeyinde birebir aynı."} | ortak,
        {"kaynak_id": "ibb_kitapcik_fatih2020", "tur": "rapor", "yil": "2020", "kunye":
            "İBB DEZİM ve BÜ KRDAE (Haziran 2020). İstanbul İli Fatih İlçesi Olası Deprem Kayıp Tahminleri Kitapçığı.",
         "yayinci": "İBB", "url": "https://depremzemin.ibb.istanbul/tr/olasi-deprem-kayip-tahminleri-ilce-kitapciklari",
         "erisim_tarihi": "kaydedilmemiş (dosya tarihi 2026-09-20)", "sha256": env["fatih.pdf"], "lisans": "telifli; dağıtılmaz",
         "yerel_ad": "arşiv: 02_kaynak_raporlar/fatih.pdf", "not": "İlçe toplamı İBB senaryo CSV'siyle birebir aynı."} | ortak,
        {"kaynak_id": "irap_istanbul", "tur": "rapor", "kunye": "İstanbul Valiliği. İstanbul İl Afet Risk Azaltma Planı (İRAP).",
         "yayinci": "İstanbul Valiliği / AFAD", "url": "", "erisim_tarihi": "kaydedilmemiş (dosya tarihi 2026-09-20)",
         "sha256": env["irap.pdf"], "lisans": "telifli; dağıtılmaz", "yerel_ad": "arşiv: 02_kaynak_raporlar/irap.pdf",
         "not": "Notlarda bu belgeden sayı alınmamış. URL bulunmalı."} | ortak,
        {"kaynak_id": "ucuncu2022", "tur": "tez", "yil": "2022", "kunye":
            "Üçüncü, Ş. (2022). Kentsel Dönüşüm Çalışmalarında Asbest Riskinin Değerlendirilmesi. Doktora tezi, Hacettepe Üniversitesi, 420 s.",
         "yayinci": "Hacettepe Üniversitesi", "url": "", "erisim_tarihi": "kaydedilmemiş (dosya tarihi 2026-09-20)",
         "sha256": env["ucuncu.pdf"], "lisans": "telifli; dağıtılmaz", "yerel_ad": "arşiv: 02_kaynak_raporlar/ucuncu.pdf",
         "not": "YÖK Tez Merkezi bağlantısı eklenmeli."} | ortak,
        {"kaynak_id": "jsmcwm2022", "tur": "rapor", "yil": "2022", "kunye":
            "Japan Society of Material Cycles and Waste Management (2022). Technical reference: Estimating the amount of disaster waste generated.",
         "yayinci": "JSMCWM", "url": "", "erisim_tarihi": "kaydedilmemiş (dosya tarihi 2026-09-20)", "sha256": env["jsmcwm2022.pdf"],
         "lisans": "telifli; dağıtılmaz", "yerel_ad": "arşiv: 02_kaynak_raporlar/jsmcwm2022.pdf", "not": "URL bulunmalı."} | ortak,
        {"kaynak_id": "epa_ap42_13_2_4", "tur": "rapor", "yil": "2006", "kunye":
            "US EPA (2006). AP-42, Section 13.2.4: Aggregate Handling and Storage Piles.", "yayinci": "US EPA", "url": "",
         "erisim_tarihi": "kaydedilmemiş (dosya tarihi 2026-09-20)", "sha256": env["handling.pdf"],
         "lisans": "ABD federal yayını (büyük olasılıkla kamu malı); kural gereği dağıtılmaz", "yerel_ad": "arşiv: 02_kaynak_raporlar/handling.pdf",
         "not": "PM emisyon faktörü için aday. URL eklenmeli."} | ortak,
        {"kaynak_id": "asbest_yonetmeligi2013", "tur": "mevzuat", "yil": "2013", "kunye":
            "Asbestle Çalışmalarda Sağlık ve Güvenlik Önlemleri Hakkında Yönetmelik. Resmî Gazete 25.01.2013, sayı 28539, Madde 11.",
         "yayinci": "T.C. Resmî Gazete", "url": "", "erisim_tarihi": "kaydedilmemiş", "sha256": "", "lisans": "mevzuat (kamu)",
         "yerel_ad": "", "okuma": "", "not": "0,1 lif/cm³ sınır değeri; Resmî Gazete bağlantısı eklenmeli, belge depoda yok."}
        | {"dagitim": "yalniz_baglanti"},
    ]


def ek_gorev_kaynaklari() -> list[dict]:
    """21 Eylül 2026 araştırma notunda (ek görevler 4.1–4.4) bağlantısı verilen kaynaklar. Hiçbiri yerelde yok."""
    ortak = {"tur": "literatur", "lisans": "yayıncı lisansı; dağıtılmaz", "dagitim": "yalniz_baglanti",
             "erisim_tarihi": "2026-09-21 (araştırma notu)", "sha256": "",
             "not": "Bağlantı araştırma notundan; belge depoda yok, değerler belgeyle karşılaştırılmadı."}
    return [
        ortak | {"kaynak_id": "ergin2024", "yil": "2024", "okuma": "tam_metin", "url": "https://dergipark.org.tr/en/download/article-file/3668294",
                 "kunye": "Ergin, Y., Demiray, G., Ergin, B. B., Ekuklu, G. (2024). 2023 Kahramanmaraş Depremi Sonrası Etkilenen Bazı "
                          "İllerde Hava Kalitesi Değerlerinin Karşılaştırılması. Afet ve Risk, 7(3), 697–712. DOI: 10.35341/afet.1421763"},
        ortak | {"kaynak_id": "vural2024", "yil": "2024", "okuma": "tam_metin",
                 "url": "https://file.techscience.com/files/RIG/2024/RIG33-1/TSP_RIG_47908/TSP_RIG_47908.pdf",
                 "kunye": "Vural, E. (2024). Assessment of Particle Matter Pollution during Post-Earthquake Debris Removal in Adiyaman City. "
                          "Revue Internationale de Géomatique, 33, 37–50. DOI: 10.32604/rig.2024.047908"},
        ortak | {"kaynak_id": "thhp_ttb2023", "yil": "2023", "okuma": "tam_metin",
                 "url": "https://temizhavahakki.org/wp-content/uploads/2025/03/THHP-TTB_ASBEST-RAPORU_Mart-2025_guncel_v2.pdf",
                 "kunye": "Yay, O. D., Yıldız, K., Gümüşel, D., Varol, G., Karababa, A. O. (Ekim 2023; Mart 2025 güncel sürüm). Deprem "
                          "Bölgesinde Üç Kentte Asbest Tespit Çalışması – Nihai Rapor. Temiz Hava Hakkı Platformu – Türk Tabipleri Birliği."},
        ortak | {"kaynak_id": "baran2026", "yil": "2026", "okuma": "ozet", "url": "https://www.researchsquare.com/article/rs-8682234/v1",
                 "kunye": "Baran, A. ve ark. (2026). Assessment of dust and fibers resulting from the demolition of buildings following an "
                          "earthquake. Research Square, v1. DOI: 10.21203/rs.3.rs-8682234/v1 (hakem değerlendirmesinden geçmemiş ön baskı)",
                 "not": "Aynı Malatya ölçümlerinin baran2024 (kongre özeti) ile ilişkisi netleştirilmeli."},
        ortak | {"kaynak_id": "diktepe2025", "yil": "2025", "okuma": "ozet", "url": "https://bulten.mta.gov.tr/makale/detay/684",
                 "kunye": "Diktepe, Ş. (Aralık 2025). İstanbul depremi senaryosunda asbest içeren inşaat ve yıkıntı atıklarının "
                          "oluşturacağı çevresel riskler. MTA Doğal Kaynaklar ve Ekonomi Bülteni."},
        ortak | {"kaynak_id": "zanoletti2024", "yil": "2024", "okuma": "tam_metin",
                 "url": "https://link.springer.com/article/10.1007/s11356-024-32592-8",
                 "kunye": "Zanoletti, A., Bontempi, E. (2024). The impacts of earthquakes on air pollution and strategies for mitigation: "
                          "a case study of Turkey. Environmental Science and Pollution Research, 31, 24662–24672."},
        ortak | {"kaynak_id": "tetik2024", "yil": "2024", "okuma": "", "url": "",
                 "kunye": "Tetik ve ark. (2024). Environmental Science and Pollution Research, 31, 9857–9866.",
                 "not": "Yalnız sunumda künye var; tam künye ve bağlantı bulunmalı."},
        ortak | {"kaynak_id": "sim_istasyon", "tur": "acik_veri", "okuma": "", "lisans": "kaydedilmemiş",
                 "url": "https://sim.csb.gov.tr/STN/STN_Report/StationDataDownloadNew",
                 "kunye": "ÇŞİDB Sürekli İzleme Merkezi, Hava Kalitesi İstasyon Veri İndirme.",
                 "not": "21.09.2026 Adıyaman sorgusunun kaydı/indirmesi yok; sayılar yeniden sorgulanmalı."},
    ]


def main() -> None:
    zaman = istek_zamanlari()
    lit = literatur()
    standart_degerler()
    lit_ek = {k["kaynak_id"]: k for k in ek_gorev_kaynaklari()}
    # oral2025 her iki notta da var: yayıncı bağlantısını öne al
    for k in lit:
        if k["kaynak_id"] == "oral2025":
            k["not"] += " Yayıncı tam metni: https://www.mdpi.com/1660-4601/22/4/649"
    tum = lit + list(lit_ek.values()) + raporlar() + acik_veri(zaman) + kadikoy_belgeleri() + belediye_belgeleri(zaman)
    ids = [k["kaynak_id"] for k in tum]
    tekrar = {i for i in ids if ids.count(i) > 1}
    if tekrar:
        raise SystemExit(f"Yinelenen kaynak_id: {sorted(tekrar)}")
    # Önceden bulunmuş arşiv bağlantılarını (src/arsiv_kontrol.py) koru
    eski = {}
    if (KOK / "kaynaklar.csv").exists():
        eski = {r["kaynak_id"]: r["arsiv_url"] for r in csv.DictReader(open(KOK / "kaynaklar.csv", encoding="utf-8"))}
    for k in tum:
        k["arsiv_url"] = k.get("arsiv_url") or eski.get(k["kaynak_id"], "")
    print("kaynaklar.csv")
    yaz(KOK / "kaynaklar.csv", tum, KAYNAK_ALANLAR)


if __name__ == "__main__":
    main()
