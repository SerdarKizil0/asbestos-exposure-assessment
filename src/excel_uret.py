"""CSV → Excel: literatür kitaplarını veri/ altındaki CSV'lerden yeniden üretir.

  python src/excel_uret.py                 ciktilar/excel/asbest_yayilim_veri_tablosu.xlsx
                                           ciktilar/excel/asbest_yayilim_standart_degerler.xlsx
  python src/excel_uret.py --karsilastir   üretilen kitapları gelen/excel/ altındaki özgünleriyle hücre hücre karşılaştırır

Tek doğruluk kaynağı CSV'dir: Excel'de düzeltme yapmayın, CSV'yi düzeltip bu betiği çalıştırın.
Değerler ve formüller birebir aynıdır; renk ve sayı biçimleri kurallarla üretilir:
  mavi yazı = kaynaktan alınan değer, siyah = formül veya metin, sarı dolgu = varsayım,
  pembe dolgu = asıl metni açılamayan kaynaktan alınmış, doğrulanmamış değer.

Diğer kitaplar kendi betikleriyle üretilir: src/tarama/rapor_xlsx.py (39 ilçe taraması),
src/tarama/kadikoy_xlsx.py (Kadıköy), model/rapor_model_xlsx.py (yayılım modeli).
"""
from __future__ import annotations

import argparse
import csv
import math
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

KOK = Path(__file__).resolve().parents[1]
VERI = KOK / "veri"
LIT = VERI / "literatur"
CIKTI = KOK / "ciktilar" / "excel"

MAVI, KOYU_MAVI, SIYAH = "FF0000FF", "FF1F4E79", "FF000000"
GRI, SARI, PEMBE, ACIK_MAVI = "FFD9D9D9", "FFFFFF00", "FFF8CBAD", "FFBDD7EE"


# ---------------------------------------------------------------- yardımcılar
def oku(yol: Path) -> list[dict]:
    with open(yol, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def sayi(v: str):
    """CSV hücresi → Excel değeri. Boş → None; sayı gibi görünen metin → int/float; formül ve metin olduğu gibi."""
    if v is None or v == "":
        return None
    if v.startswith("="):
        return v
    try:
        if v.lstrip("-").isdigit():
            return int(v)
        f = float(v)
        return f if math.isfinite(f) else v
    except ValueError:
        return v


def metin(v: str):
    return None if v in (None, "") else v


def dolgu(renk: str) -> PatternFill:
    return PatternFill("solid", fgColor=renk)


def yaz(ws, satir: int, sutun: int, deger, *, renk=SIYAH, kalin=False, italik=False, boyut=10, zemin=None,
        bicim=None, sar=True):
    c = ws.cell(satir, sutun, deger)
    c.font = Font(name="Calibri", size=boyut, bold=kalin, italic=italik, color=renk)
    c.alignment = Alignment(wrap_text=sar, vertical="top")
    if zemin:
        c.fill = dolgu(zemin)
    if bicim:
        c.number_format = bicim
    return c


def baslik_satiri(ws, satir: int, basliklar: list[str]) -> None:
    for i, b in enumerate(basliklar, 1):
        yaz(ws, satir, i, b, kalin=True, zemin=GRI)


def genislik(ws, gen: dict) -> None:
    for harf, g in gen.items():
        ws.column_dimensions[harf].width = g


def sayi_bicimi(v, birim: str = "") -> str | None:
    if not isinstance(v, (int, float)) or isinstance(v, bool):
        return None
    if "oran" in (birim or "") or "yüzde" in (birim or ""):
        return "0.0%" if abs(round(v * 100, 1) - v * 100) < 1e-9 else "0.00%"
    if abs(v) >= 1000 and float(v).is_integer():
        return "#,##0"
    return None


# ---------------------------------------------------------------- veri tablosu kitabı
def kaynak_tablosu():
    kay = {k["kaynak_id"]: k for k in oku(KOK / "kaynaklar.csv")}
    roller = oku(LIT / "kaynak_rolleri.csv")
    return kay, roller


def ozet(wb, n_kaynak: int) -> None:
    ws = wb.active
    ws.title = "Özet"
    yaz(ws, 1, 1, "Deprem ve bina çöküşü sonrası asbest yayılımı ", kalin=True, boyut=13, sar=False)
    baslik_satiri(ws, 4, ["Kaynak sayımı", "Sayı", "Açıklama"])
    son = n_kaynak + 1
    rng = f"Kaynaklar!$E$2:$E${son}"
    satirlar = [
        ("Toplam kaynak", f"=COUNTA(Kaynaklar!$A$2:$A${son})", "Belgedeki kaynakçanın tamamı."),
        ("Model parametresi veren", f'=COUNTIF({rng},"Model parametresi")', "Salım faktörü, dönüşüm katsayısı, lif fiziği, kaynak bileşimi."),
        ("Model kurulumu veren", f'=COUNTIF({rng},"Model kurulumu")', "Alan, ızgara, kaynak geometrisi, meteoroloji."),
        ("Doğrulama ölçümü veren", f'=COUNTIF({rng},"Doğrulama ölçümü")', "Modeli karşılaştırmak için sahada ölçülmüş değerler."),
        ("Model sapması veren", f'=COUNTIF({rng},"Model sapması")', "Modelin ölçüme göre düşük/yüksek tahmin katsayıları."),
        ("Doğrulanamayan (asıl metin açılamadı)", f'=COUNTIF({rng},"Doğrulanamadı")', "Değerleri girildi ama pembe işaretli."),
        ("Sayısal yayılım verisi olmayan", f'=COUNTIF({rng},"Sayısal yayılım verisi yok")', "Niteliksel bulgu veya yalnızca tespit sayısı."),
        ("Sayısal yayılım verisi olan", "=B5-B11", "Toplam eksi sayısal verisi olmayanlar."),
        ("Bunlardan doğrulanmış", "=B12-B10", "Sayısal verisi olan eksi doğrulanamayanlar."),
    ]
    for i, (a, b, c) in enumerate(satirlar, 5):
        yaz(ws, i, 1, a)
        yaz(ws, i, 2, b, bicim="0")
        yaz(ws, i, 3, c)
    baslik_satiri(ws, 15, ["Sayfa", None, "İçerik"])
    sayfalar = [
        ("Kaynaklar", f"{n_kaynak} kaynak: künye, rol, doğrulama durumu ve bağlantı."),
        ("Parametreler", "Modele doğrudan girebilecek sabitler: kaynak terimi, salım, dönüşüm, lif fiziği, model belirsizliği."),
        ("Ölçümler", "Doğrulama için sahada ölçülmüş değerler (lif ve toz derişimleri, mesafeler, süreler)."),
        ("Model_Kurulumları", "Dört modelleme çalışmasının kurulumu yan yana."),
        ("Model_Sonuçları", "WTC CALPUFF seyrelme–mesafe–meteoroloji sonuçları ve Kobe hesap–ölçüm karşılaştırması."),
        ("Hesaplar", "Makalelerdeki hesapların formüllerle yeniden yapılması, birim dönüşümü ve çökelme hızı."),
    ]
    for i, (a, c) in enumerate(sayfalar, 16):
        yaz(ws, i, 1, a, kalin=True)
        yaz(ws, i, 3, c)
    baslik_satiri(ws, 23, ["Renk ve biçim açıklaması", None, None])
    yaz(ws, 24, 1, "Mavi yazı", renk=MAVI); yaz(ws, 24, 3, "Kaynaktan alınan değer (elle girilmiş).")
    yaz(ws, 25, 1, "Siyah yazı"); yaz(ws, 25, 3, "Formül veya metin.")
    yaz(ws, 26, 1, "Sarı dolgu", zemin=SARI); yaz(ws, 26, 3, "Varsayım ya da değiştirilebilir girdi (kaynakta yok).")
    yaz(ws, 27, 1, "Pembe dolgu", zemin=PEMBE); yaz(ws, 27, 3, "Asıl metni açılamayan kaynaktan alınmış, doğrulanmamış değer.")
    yaz(ws, 29, 1, "Notlar", kalin=True)
    notlar = [
        "Oran sütunları kesir olarak tutulur, yüzde biçiminde gösterilir (ör. 0,0001 = %0,01).",
        "Aralık verilen değerler 'Alt' ve 'Üst' sütunlarına, tek değerler 'Değer' sütununa yazıldı. 'Değer' boşsa kaynak yalnızca aralık veya sınır vermiştir.",
        "PCM optik mikroskop, TEM ve SEM elektron mikroskobu sayımıdır; aynı örnekte sonuçlar farklı çıkabilir (bkz. Ölçümler M18–M20).",
        "Boş hücre, kaynakta o bilgi için sayı verilmediği anlamına gelir.",
        "Hesaplar sayfasındaki 'Fark' sütunu, formülle yeniden hesaplanan değerin makaledeki değerden yüzde farkıdır.",
    ]
    for i, n in enumerate(notlar, 30):
        yaz(ws, i, 1, n)
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=3)
        ws.row_dimensions[i].height = 27.75
    genislik(ws, {"A": 58, "B": 14, "C": 70})


def kaynaklar_sayfasi(wb, kay: dict, roller: list[dict]) -> set:
    ws = wb.create_sheet("Kaynaklar")
    baslik_satiri(ws, 1, ["No", "Kısa ad", "Künye", "Olay / konu", "Birincil rol", "Diğer roller / not",
                          "Excel'deki sayfalar", "Doğrulama durumu", "Bağlantı"])
    pembe = set()
    for i, r in enumerate(roller, 2):
        k = kay[r["kaynak_id"]]
        zemin = PEMBE if r["birincil_rol"] == "Doğrulanamadı" else None
        if zemin:
            pembe.add(r["kaynak_id"])
        degerler = [sayi(r["no"]), r["kisa_ad"], k["kunye"], r["olay_konu"], r["birincil_rol"], r["diger_roller"],
                    r["excel_sayfalari"], r["okuma_durumu"], metin(k["url"])]
        for j, v in enumerate(degerler, 1):
            yaz(ws, i, j, v, renk=KOYU_MAVI if j == 9 else SIYAH, zemin=zemin)
    ws.freeze_panes = "C2"
    genislik(ws, {"A": 5, "B": 24, "C": 60, "D": 26, "E": 22, "F": 30, "G": 32, "H": 26, "I": 50})
    return pembe


def tablo_sayfasi(wb, ad: str, basliklar: list[str], satirlar: list[list], *, sayisal: set[int], birim_sutunu=None,
                  pembe_satir=None, dondur=None, gen=None) -> None:
    ws = wb.create_sheet(ad)
    baslik_satiri(ws, 1, basliklar)
    for i, r in enumerate(satirlar, 2):
        zemin = PEMBE if pembe_satir and pembe_satir(r) else None
        birim = r[birim_sutunu] if birim_sutunu is not None else ""
        for j, v in enumerate(r, 1):
            mavi = (j - 1) in sayisal and isinstance(v, (int, float))
            yaz(ws, i, j, v, renk=MAVI if mavi else SIYAH, zemin=zemin, bicim=sayi_bicimi(v, birim) if mavi else None)
    if dondur:
        ws.freeze_panes = dondur
    genislik(ws, gen or {})


def model_sonuclari(wb) -> None:
    ws = wb.create_sheet("Model_Sonuçları")
    yaz(ws, 1, 1, "A) WTC 2001 – CALPUFF seyrelme sonuçları (Gilliam ve ark. 2005). Seyrelme = kaynak hacmindeki derişim / "
                  "yer düzeyi derişimi. Model 2 km'den yakında geçersiz.", kalin=True, sar=False)
    baslik_satiri(ws, 2, ["Tarih / dönem", "Meteorolojik durum", "Karışım yük. alt (m)", "Karışım yük. üst (m)",
                          "Rüzgâr alt (m/s)", "Rüzgâr üst (m/s)", "Seyrelme alt", "Seyrelme üst", "Mesafe alt (km)",
                          "Mesafe üst (km)", "Gözlem (PM2.5)", "Not"])
    alan = ["tarih_donem", "meteorolojik_durum", "karisim_alt_m", "karisim_ust_m", "ruzgar_alt_ms", "ruzgar_ust_ms",
            "seyrelme_alt", "seyrelme_ust", "mesafe_alt_km", "mesafe_ust_km", "gozlem_pm25", "not"]
    a = oku(LIT / "wtc_seyrelme.csv")
    for i, r in enumerate(a, 3):
        for j, k in enumerate(alan, 1):
            v = sayi(r[k]) if 3 <= j <= 10 else metin(r[k])
            yaz(ws, i, j, v, renk=MAVI if isinstance(v, (int, float)) else SIYAH, bicim=sayi_bicimi(v))
    n = 3 + len(a)
    yaz(ws, n, 1, "Kaynak: Kaynak No 2 (Gilliam ve ark. 2005), §3.1–3.2 ve §4. Boş hücre = makalede o koşul için sayı "
                  "verilmemiş.", italik=True, sar=False)
    yaz(ws, n + 2, 1, "B) Kobe 1995 – plume-puff hesap ile ölçüm karşılaştırması (Terazono ve ark. 2000, §4.3.3–4.3.4)",
        kalin=True, sar=False)
    baslik_satiri(ws, n + 3, ["Kalem", "Açıklama", "Değer", "Alt", "Üst", "Birim"])
    b = oku(LIT / "kobe_karsilastirma.csv")
    for i, r in enumerate(b, n + 4):
        for j, k in enumerate(["kalem", "aciklama", "deger", "alt", "ust", "birim"], 1):
            v = sayi(r[k]) if 3 <= j <= 5 else metin(r[k])
            yaz(ws, i, j, v, renk=MAVI if isinstance(v, (int, float)) else SIYAH)
    yaz(ws, n + 4 + len(b), 1, "Kaynak: Kaynak No 1 (Terazono ve ark. 2000).", italik=True, sar=False)
    ws.freeze_panes = "B3"
    genislik(ws, {"A": 24, "B": 34, "C": 12, "E": 11, "G": 12, "I": 11, "K": 44, "L": 50})


def hesaplar(wb) -> None:
    ws = wb.create_sheet("Hesaplar")
    for r in oku(LIT / "hesaplar.csv"):
        s = int(r["satir"])
        tur = r["tur"]
        hucreler = [metin(r["kalem"]), sayi(r["deger"]), metin(r["birim"]), sayi(r["makale_degeri"]), sayi(r["fark"]),
                    metin(r["not"])]
        if tur == "baslik":
            yaz(ws, s, 1, hucreler[0], sar=False)
            continue
        zemin = {"bolum": ACIK_MAVI, "sutun_basligi": GRI}.get(tur)
        for j, v in enumerate(hucreler, 1):
            renk, z, bicim = SIYAH, zemin, None
            if tur == "kalem":
                if j == 2:
                    bicim = r["deger_bicimi"]
                    if r["girdi"] == "kaynak":
                        renk = MAVI
                    elif r["girdi"] == "varsayim":
                        z = SARI
                elif j == 4:
                    bicim = r["makale_bicimi"]
                    renk = MAVI if isinstance(v, (int, float)) else SIYAH
                elif j == 5 and v is not None:
                    bicim = "0.0%"
            if v is not None or z:
                yaz(ws, s, j, v, renk=renk, kalin=tur in ("bolum", "sutun_basligi"), zemin=z, bicim=bicim)
    genislik(ws, {"A": 52, "B": 18, "C": 16, "D": 18, "E": 12, "F": 56})


def veri_tablosu() -> Path:
    kay, roller = kaynak_tablosu()
    wb = openpyxl.Workbook()
    ozet(wb, len(roller))
    pembe = kaynaklar_sayfasi(wb, kay, roller)

    par = oku(VERI / "parametreler.csv")
    tablo_sayfasi(
        wb, "Parametreler",
        ["ID", "Grup", "Parametre", "Değer", "Alt", "Üst", "Birim", "Koşul / kapsam", "Kaynak No", "Kaynak", "Yer", "Not"],
        [[r["parametre_id"], r["grup"], r["parametre"], sayi(r["deger"]), sayi(r["alt"]), sayi(r["ust"]), r["birim"],
          metin(r["kosul"]), sayi(r["kaynak_no"]), r["kaynak"], metin(r["yer"]), metin(r["not"])] for r in par],
        sayisal={3, 4, 5}, birim_sutunu=6, pembe_satir=lambda r: r[8] in {20, 21}, dondur="D2",
        gen={"A": 6, "B": 16, "C": 44, "D": 14, "E": 12, "G": 18, "H": 44, "I": 9, "J": 24, "K": 18, "L": 48})
    assert {20, 21} == {int(r["kaynak_no"]) for r in par if r["kaynak_id"] in pembe}

    olc = oku(LIT / "olcumler.csv")
    tablo_sayfasi(
        wb, "Ölçümler",
        ["ID", "Olay", "Yer / nokta", "Dönem", "Ölçülen", "Yöntem", "Değer", "Alt", "Üst", "Birim", "N (örnek/nokta)",
         "Kaynak No", "Kaynak", "Yer", "Not"],
        [[r["olcum_id"], r["olay"], r["yer_nokta"], r["donem"], r["olculen"], r["yontem"], sayi(r["deger"]), sayi(r["alt"]),
          sayi(r["ust"]), r["birim"], sayi(r["n"]), sayi(r["kaynak_no"]), r["kaynak"], metin(r["yer"]), metin(r["not"])]
         for r in olc],
        sayisal={6, 7, 8}, birim_sutunu=9, dondur="E2",
        gen={"A": 6, "B": 18, "C": 40, "D": 20, "E": 34, "F": 22, "G": 11, "H": 10, "J": 11, "K": 14, "L": 9, "M": 24})

    mk = oku(LIT / "model_kurulumlari.csv")
    alan = ["calisma", "olay", "model", "alan_olcek", "yatay_cozunurluk", "zaman_cozunurlugu", "kaynak_tanimi", "salim",
            "meteoroloji_girdisi", "cokelme", "gecerlilik_kisit", "dogrulama_verisi", "yer", "kaynak_no"]
    tablo_sayfasi(
        wb, "Model_Kurulumları",
        ["Çalışma", "Olay", "Model", "Alan / ölçek", "Yatay çözünürlük", "Zaman çözünürlüğü", "Kaynak tanımı", "Salım",
         "Meteoroloji girdisi", "Çökelme", "Geçerlilik / kısıt", "Doğrulama verisi", "Yer", "Kaynak No"],
        [[sayi(r[k]) if k == "kaynak_no" else metin(r[k]) for k in alan] for r in mk],
        sayisal={13}, dondur="B2",
        gen={"A": 22, "B": 16, "C": 30, "D": 26, "E": 22, "G": 36, "J": 24, "K": 40, "L": 28, "M": 14, "N": 9})

    model_sonuclari(wb)
    hesaplar(wb)
    wb.properties.creator = "asbestos-exposure-assessment (src/excel_uret.py)"
    CIKTI.mkdir(parents=True, exist_ok=True)
    yol = CIKTI / "asbest_yayilim_veri_tablosu.xlsx"
    wb.save(yol)
    return yol


# ---------------------------------------------------------------- standart değerler kitabı
KOYU, BEYAZ = "FF404040", "FFFFFFFF"
BOLUM_ZEMIN = ["FFDDEBF7", "FFE2EFDA", "FFFCE4D6"]


def standart_degerler() -> Path:
    kay = {k["kaynak_id"]: k for k in oku(KOK / "kaynaklar.csv")}
    satirlar = oku(LIT / "standart_degerler.csv")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Standart Değerler"
    yaz(ws, 1, 1, "Asbest Yayılım Modeli – Standart Değerler", renk=None, kalin=True, boyut=15, sar=False)
    ws.row_dimensions[1].height = 26
    for j, b in enumerate(["Parametre", "Standart değer", "Birim", "Kaynak"], 1):
        yaz(ws, 4, j, b, renk=BEYAZ, kalin=True, boyut=11, zemin=KOYU, sar=False)
    ws.row_dimensions[4].height = 22
    s, bolum, nb = 4, None, -1
    for r in satirlar:
        if r["bolum"] != bolum:
            s += 1
            nb += 1
            bolum = r["bolum"]
            yaz(ws, s, 1, bolum, renk=None, kalin=True, boyut=11, zemin=BOLUM_ZEMIN[nb], sar=False)
            ws.merge_cells(start_row=s, start_column=1, end_row=s, end_column=4)
            ws.row_dimensions[s].height = 22
        s += 1
        for j, (v, kalin) in enumerate([(r["parametre"], False), (r["gosterim"], True), (r["birim"], False),
                                        (r["kaynak"], False)], 1):
            yaz(ws, s, j, v, renk=None, kalin=kalin, boyut=11)
        ws.row_dimensions[s].height = 20
    s += 2
    yaz(ws, s, 1, "Açıklamalar", renk=BEYAZ, kalin=True, boyut=11, zemin=KOYU, sar=False)
    goruldu = []
    for r in satirlar:
        if r["aciklama_basligi"] in goruldu:
            continue
        goruldu.append(r["aciklama_basligi"])
        s += 1
        yaz(ws, s, 1, r["aciklama_basligi"], renk=None, kalin=True)
        yaz(ws, s, 2, r["aciklama"], renk=None)
        ws.merge_cells(start_row=s, start_column=2, end_row=s, end_column=4)
    s += 2
    yaz(ws, s, 1, "Birim: lif/L = 1 litre havadaki lif sayısı. 1 lif/cm³ = 1.000 lif/L.", renk="FF595959", sar=False)
    ws.freeze_panes = "A5"
    genislik(ws, {"A": 54, "B": 20, "C": 24, "D": 38})

    ws = wb.create_sheet("Kaynaklar")
    for j, b in enumerate(["Kaynak", "Künye"], 1):
        yaz(ws, 1, j, b, renk=BEYAZ, kalin=True, boyut=11, zemin=KOYU, sar=False)
    for i, r in enumerate(oku(LIT / "standart_degerler_kaynaklari.csv"), 2):
        kunye = kay[r["kaynak_id"]]["kunye"] + (" " + r["ek_not"] if r["ek_not"] else "")
        yaz(ws, i, 1, r["kisa_ad"], renk=None, kalin=True)
        yaz(ws, i, 2, kunye, renk=KOYU_MAVI)
    genislik(ws, {"A": 30, "B": 110})
    wb.properties.creator = "asbestos-exposure-assessment (src/excel_uret.py)"
    yol = CIKTI / "asbest_yayilim_standart_degerler.xlsx"
    wb.save(yol)
    return yol


# ---------------------------------------------------------------- karşılaştırma
def karsilastir(uretilen: Path, ozgun: Path) -> int:
    a, b = openpyxl.load_workbook(uretilen), openpyxl.load_workbook(ozgun)
    fark = 0
    if a.sheetnames != b.sheetnames:
        print(f"  sayfa adları farklı: {a.sheetnames} ≠ {b.sheetnames}")
        fark += 1
    for ad in b.sheetnames:
        wa, wb_ = a[ad], b[ad]
        hucre = {}
        for ws, i in ((wa, 0), (wb_, 1)):
            for row in ws.iter_rows():
                for c in row:
                    if c.value is not None:
                        hucre.setdefault(c.coordinate, [None, None])[i] = c.value
        for k, (x, y) in sorted(hucre.items()):
            esit = x == y or (isinstance(x, (int, float)) and isinstance(y, (int, float)) and math.isclose(x, y, rel_tol=1e-12))
            if not esit:
                fark += 1
                if fark <= 15:
                    print(f"  {ad}!{k}: üretilen={x!r:.80} özgün={y!r:.80}")
    print(f"  {uretilen.name}: {fark} farklı hücre")
    return fark


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--karsilastir", action="store_true", help="gelen/excel altındaki özgün kitaplarla karşılaştır")
    arg = ap.parse_args()
    yollar = [veri_tablosu(), standart_degerler()]
    for y in yollar:
        print(f"yazıldı: {y.relative_to(KOK)}")
    if arg.karsilastir:
        toplam = sum(karsilastir(y, KOK / "gelen" / "excel" / y.name) for y in yollar)
        sys.exit(1 if toplam else 0)


if __name__ == "__main__":
    main()
