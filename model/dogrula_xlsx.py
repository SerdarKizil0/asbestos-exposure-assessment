"""Excel modelinin yeniden hesaplanmış değerlerini Python sonuçlarıyla karşılaştırır.

Önce: python3 rapor_model_xlsx.py && python3 /mnt/skills/public/xlsx/scripts/recalc.py cikti/kadikoy_asbest_yayilim_modeli.xlsx
(ya da dosyayı Excel/LibreOffice'te açıp kaydedin; formül değerleri ancak o zaman dosyaya yazılır).
"""
import json
import sys

from openpyxl import load_workbook

from ortak import CIKTI

S = json.loads((CIKTI / "sonuclar.json").read_text())
wb = load_workbook(CIKTI / "kadikoy_asbest_yayilim_modeli.xlsx", data_only=True)
sd = {d["mahalle"]: d for d in S["mahalle"]}
ws = wb["Sonuçlar"]
if ws["I2"].value is None:
    sys.exit("Formül değerleri yok: dosyayı önce yeniden hesaplatın.")
mx = 0.0
for r in range(2, 2 + len(sd)):
    a = ws.cell(r, 1).value
    for col, k in [(9, "c_calisma"), (15, "c_kis"), (16, "c_yaz"), (17, "c_S1"), (18, "c_S2"), (19, "c_S3"), (20, "c_S4")]:
        mx = max(mx, abs(ws.cell(r, col).value - sd[a][k]) / sd[a][k])
    assert ws.cell(r, 13).value == sd[a]["sira"], f"sıra farkı: {a}"
print(f"Mahalle derişimleri: en büyük göreli fark {mx:.1e}; sıralar aynı")
pwm = next(ws.cell(r, 9).value for r in range(20, 40) if str(ws.cell(r, 1).value).startswith("Nüfus ağırlıklı ortalama, çalışma"))
print(f"Nüfus ağırlıklı ortalama: Excel {pwm:.6f}, Python {S['pwm_calisma']:.6f}")
ws = wb["Tek yıkım"]
mes = S["tek_yikim"]["mesafeler"]
satirlar = [r for r in range(24, 60) if ws.cell(r, 2).value is not None and isinstance(ws.cell(r, 4).value, float)]
mxf = 0.0
for r, t in zip(satirlar, S["tek_yikim"]["tablo"]):
    ex = [ws.cell(r, 4 + j).value for j in range(len(mes))]
    py = [t[f"{m}_m"] for m in mes]
    mxf = max(mxf, max(abs(e - p) / p for e, p in zip(ex, py)))
print(f"Tek yıkım ({len(satirlar)} satır): en büyük göreli fark %{100 * mxf:.2f}")
assert mx < 1e-6 and mxf < 0.01 and len(satirlar) == len(S["tek_yikim"]["tablo"])
print("Tamam")
