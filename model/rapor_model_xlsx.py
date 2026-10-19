"""Canlı Excel modeli → model/cikti/kadikoy_asbest_yayilim_modeli.xlsx

Renk düzeni kullanıcının kendi literatür tablosuyla aynı:
  mavi yazı = kaynaktan alınan değer; siyah yazı = formül ya da metin; sarı dolgu = varsayım / değiştirilebilir girdi;
  pembe dolgu = asıl metni açılamamış kaynaktan, doğrulanmamış değer; gri yazı = Python modelinin hesapladığı sabit değer.
Parametreler sayfasındaki sarı/mavi hücreler değişince salım, derişim, sıra ve maruziyet sütunları yeniden hesaplanır.
Yayılım katsayıları (D_* sayfaları) ve duyarlılık/Monte Carlo tabloları Python çıktısıdır; Excel'de yeniden hesaplanmaz.
"""
from __future__ import annotations

import json

import numpy as np
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from girdiler import kadikoy_asbest, kur as mahalle_tablosu
from model import CALISMA_PAYI, CALISMA_SAATI_AY, P, aktarimlar
from ortak import CIKTI, GIRDI, RAW

F = "Arial"
MAVI = Font(name=F, size=10, color="0000FF")
SIYAH = Font(name=F, size=10, color="000000")
GRI = Font(name=F, size=10, color="7F7F7F")
KALIN = Font(name=F, size=10, bold=True)
BASLIK = Font(name=F, size=12, bold=True)
BEYAZ = Font(name=F, size=10, bold=True, color="FFFFFF")
SARI = PatternFill("solid", fgColor="FFFF00")
PEMBE = PatternFill("solid", fgColor="F8CBAD")
HDR = PatternFill("solid", fgColor="404040")
ACIK = PatternFill("solid", fgColor="F2F2F2")
SAR = Alignment(wrap_text=True, vertical="top")

S = json.loads((CIKTI / "sonuclar.json").read_text())
DU = json.loads((CIKTI / "duyarlilik.json").read_text())
MO = json.loads((GIRDI / "meteo_ozet.json").read_text())


def baslik_satiri(ws, satir, basliklar, genislik=None):
    for j, b in enumerate(basliklar, 1):
        c = ws.cell(row=satir, column=j, value=b)
        c.font, c.fill, c.alignment = BEYAZ, HDR, Alignment(wrap_text=True, vertical="top")
    if genislik:
        for j, w in enumerate(genislik, 1):
            ws.column_dimensions[get_column_letter(j)].width = w


def yaz(ws, r, c, v, font=SIYAH, fill=None, fmt=None, wrap=False):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    if wrap:
        cell.alignment = SAR
    return cell


def ad_tanimla(wb, ad, ref):
    wb.defined_names[ad] = DefinedName(ad, attr_text=ref)


# ----------------------------------------------------------------------------------------------------------------
def okuma_notu(ws):
    s = S
    top = sorted(s["mahalle"], key=lambda d: d["sira"])
    satirlar = [
        ("Kadıköy: Mw 7,5 deprem sonrası yıkım ve enkaz kaldırma kaynaklı asbest lifi yayılımı — tarama modeli (14 Ekim 2026)", BASLIK),
        ("", None),
        ("Bu kitap ne yapar", KALIN),
        ("İBB/KOERİ 2019 deprem senaryosundaki (Mw 7,5) mahalle bazında ağır ve çok ağır hasarlı binaların yıkılacağını varsayar; Kadıköy Belediyesi'nin 2015-2025 asbest verisinden "
         "binada asbest bulunma olasılığını ve bina başına asbestli malzeme miktarını alır; literatür tablonuzdaki salım faktörü ve lif dönüşümüyle havaya geçen lif sayısını hesaplar; "
         "10 yıllık saatlik ERA5 rüzgâr, kararlılık ve karışım yüksekliğiyle (çalışma saatleri) Gauss tipi yayılım modeli çalıştırır ve her Kadıköy mahallesindeki ortalama lif artışını verir.", None),
        ("", None),
        ("Nasıl kullanılır", KALIN),
        ("Parametreler sayfasındaki sarı ve mavi hücreleri değiştirin: Kaynak mahalleler (salım) ve Sonuçlar (derişim, sıra, maruziyet) sayfaları kendiliğinden yeniden hesaplanır. "
         "Yayılım katsayıları (D_ sayfaları), Duyarlılık ve Monte Carlo tabloları Python modelinin çıktısıdır (gri yazı); bunlar Excel'de değişmez, model/ klasöründeki betikle yeniden üretilir.", None),
        ("", None),
        ("Temel durumdaki sonuçlar (mutlak değerler çok belirsiz; sıralama sağlam)", KALIN),
        (f"İlçe geneli nüfus ağırlıklı ortalama artış (çalışma saatleri): {s['pwm_calisma']:.4f} lif/L. Karşılaştırma: olağan kent havası 0,19-0,41 lif/L, Kobe 1995 deprem sonrası 1-2 lif/L. "
         f"Monte Carlo (4.000 deneme) %90 aralığı: {DU['mc']['pwm_yuzdelik']['5']:.4f}-{DU['mc']['pwm_yuzdelik']['95']:.2f} lif/L.", None),
        ("En yüksek beş mahalle: " + ", ".join(f"{d['mahalle']} (%{100*d['mc_ilk5']:.0f})" for d in top[:5]) +
         " — parantezde Monte Carlo'da ilk beşte çıkma olasılığı. En düşükler: " + ", ".join(d["mahalle"] for d in top[-5:]) + ".", None),
        ("Yıkımlar durgun ya da kararlı sabah havasına denk gelirse ilçe ortalaması yaklaşık 4,7 katına, en yüksek mahalle değeri 5-6 katına çıkıyor (Fikirtepe'de 0,0035'ten 0,022 lif/L'ye). "
         "Poyraz (çalışma saatlerinin %56'sı) dumanı güneybatıya, kıyıya iter: Caferağa, Feneryolu, Fenerbahçe, Zühtüpaşa göreli olarak artar. "
         "Lodos (%22, kışın %30) kuzeydoğuya: Hasanpaşa, Dumlupınar, Acıbadem, Koşuyolu. Güneyden gelen kararlı sabah akışında Acıbadem ve Koşuyolu 9-10 katına çıkar.", None),
        ("Asıl yüksek derişim yıkım sahasının yakınında: asbestli tek bir binanın ıslatmasız yıkımında rüzgâr altı 50 m'de 1,8-10 lif/L, 100 m'de 0,6-5 lif/L (temel salım); "
         "gevrek malzeme ya da kuru kırmada bunun 10 katı (20 m'de 36-196 lif/L). Kobe'de sökümsüz yıkımın parsel sınırında 160-250 lif/L ölçülmüş.", None),
        ("", None),
        ("Sonucu en çok değiştirenler", KALIN),
        ("'Ne kadar?' sorusunu salım faktörü (varyansın %46'sı), kütleden lif sayısına dönüşüm (%12) ve bina başına asbestli malzeme (%10) belirliyor; meteoroloji ve yayılım seçenekleri ilçe ortalamasında %1'in altında.", None),
        ("'Nerede?' sorusunu neredeyse yalnız binaların nerede yıkılacağı belirliyor: yıkımlar deprem senaryosu yerine kentsel dönüşümdeki riskli yapı dağılımını izlerse sıralama tersine dönüyor "
         "(Spearman ρ = -0,80; ilk beş Suadiye, Erenköy, Göztepe, Caddebostan, Bostancı olur).", None),
        ("", None),
        ("Renk açıklaması (sizin tablonuzla aynı düzen)", KALIN),
        ("Mavi yazı: kaynaktan alınan değer (elle girilmiş).", MAVI),
        ("Siyah yazı: formül ya da metin.", SIYAH),
        ("Sarı dolgu: varsayım ya da değiştirilebilir girdi (kaynakta yok ya da aralık seçimi bize ait).", None),
        ("Pembe dolgu: asıl metni açılamamış kaynaktan alınmış, doğrulanmamış değer.", None),
        ("Gri yazı: Python modelinin hesapladığı sabit değer (Excel'de yeniden hesaplanmaz).", GRI),
        ("", None),
        ("Sınırlar", KALIN),
        ("Tarama modelidir; İstanbul'da yapılmış bir asbest ölçümüyle doğrulanmadı. Bina gövdeleri, sokak kanyonları, deniz meltemi, yağışla bastırma, çöken binanın ilk toz bulutu, "
         "enkazın taşınması ve döküm sahaları modelde yok. Kaynaklar mahalle içinde eşit dağıtıldı (bina konumu yok). ERA5 kıyı hücresi durgun saatleri olduğundan az gösterebilir. "
         "Gauss modelleri 200 m altında ölçümü 1,6-2,6 kat düşük tahmin edebilir (Duke Forest 1996). Lif sayısı PCM eşdeğeri varsayıldı; sağlık riskine çevirmek için ayrı hesap gerekir.", None),
    ]
    ws.column_dimensions["A"].width = 125
    for i, (t, f) in enumerate(satirlar, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = f if isinstance(f, Font) else SIYAH
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if "Sarı dolgu" in t:
            c.fill = SARI
        if "Pembe dolgu" in t:
            c.fill = PEMBE


def tablolariniz(ws):
    baslik_satiri(ws, 1, ["Dosya", "Sayfa / hücre", "Ne var", "Modelde", "Not"], [16, 22, 46, 16, 90])
    satirlar = [
        ("Literatür", "Parametreler P20", "Salım faktörü %0,01 (sökümsüz yıkım; EPA faktörü, Terazono 2000)", "Kullanıldı (temel)", "Aralık %0,001-%0,1 olarak genişletildi; asbestli çimento (bağlı) malzeme için doğrudan ölçüm yok. Sonucu en çok değiştiren parametre."),
        ("Literatür", "Parametreler P26", "Kütleden lif sayısına dönüşüm 30 lif/ng (Terazono 2000, Şekil 3)", "Kullanıldı (temel)", "Hangi sayım yöntemine (PCM/TEM) ait olduğu kontrol edilmeli; Ölçümler M18-M20'de aynı filtre PCM 250, TEM 5.300 lif/L. Aralık 10-100 alındı."),
        ("Literatür", "Model_Sonuçları C25:C26, B27", "Kobe'de ölçülen artış hesaplananın yaklaşık 5 katı (eğim 4,94; R² 0,41)", "Duyarlılıkta (×5)", "Salım zincirinin düşük tahmin ettiğini gösteriyor; temel durumda düzeltme yok, duyarlılıkta ×1-×5."),
        ("Literatür", "Parametreler P01", "Püskürtme asbest yükü 0,086 kg/m² (Japonya, 1957-1980 S/RC/SRC)", "Kullanılmadı", "Japonya'daki gevrek püskürtme asbest ve tüm kat alanı için. Yerine Kadıköy'ün ölçülmüş değeri kullanıldı: asbestli bina başına 1.389 kg asbestli malzeme (2015-2025 havuz), çoğu asbestli çimento."),
        ("Literatür", "Parametreler P04", "Hasarlı binalarda püskürtme asbest oranı %3-20 (Kobe)", "Kullanılmadı", "Yerine Kadıköy'ün tespit oranı: %19,4 (havuz), yıllara göre %2,4-%36,3."),
        ("Literatür", "Parametreler P08, P21", "Yıkım oranı %20 (3 ay); gündüz düzeltmesi ×2,5", "Uyarlandı", "Model yalnız çalışma saatlerinde (08-18, pzt-cmt) salım yapıyor; gündüz düzeltmesine gerek yok. Süre ayrı parametre (6-24 ay)."),
        ("Literatür", "Parametreler P27-P33", "Lif aerodinamiği, çökelme hızı, çökelmenin payı %1-2", "Gerekçe", "Kilometre ölçeğinde çökelmenin ihmal edilebileceğini destekliyor; modelde çökelme yok."),
        ("Literatür", "Parametreler P34-P38", "Gauss modellerinin yakın alanda 1,6-2,6 kat düşük tahmini; kutu modelinin ≥10 kat aşırı tahmini", "Sınır notu", "Yakın alan (Tek yıkım) sonuçları için belirsizlik notu olarak kullanıldı."),
        ("Literatür", "Ölçümler M03-M04, M01-M02, M17-M18", "Olağan ortam 0,19-0,41 lif/L; Kobe deprem sonrası 1-2 (en çok 6); sökümsüz yıkım sınırı 160-250 lif/L", "Karşılaştırma", "Model gevrek malzeme + durgun havada 20 m'de 150-200 lif/L veriyor: Kobe sınır ölçümüyle aynı mertebe."),
        ("Literatür", "Ölçümler M34-M36", "Malatya 2024 lifli toz (PCM): aktif yıkım 4-10, döküm sahası 12, kontrol 1 lif/L; SEM/EDXA'da asbest lifi saptanmamış", "Karşılaştırma", "Yakın alan değerleriyle aynı mertebe; asbest teyidi olmayan PCM sayımları olduğu için doğrulama sayılmadı."),
        ("Literatür", "Model_Sonuçları A (WTC CALPUFF)", "Seyrelme katsayıları, 2 km'den uzakta geçerli", "Niteliksel", "Tek dev kaynak (500×500×50 m) için; İstanbul'da yüzlerce küçük yıkım var. Durgun gece = yüksek derişim bulgusu modelle uyumlu."),
        ("Literatür", "Parametreler P24-P25 (pembe)", "Eternit levhadan salım (doğrulanmamış)", "Kullanılmadı", "Asıl metin açılmadığı için."),
        ("Literatür", "Tümü", "İstanbul'a ait hasar, bina, nüfus, meteoroloji verisi", "EKSİKTİ", "Bu kitapta eklendi: İBB deprem senaryosu, Kadıköy açık verisi (nüfus, 65+, bina yaşı, toplanma alanı, sınırlar), ERA5 2015-2024."),
        ("Literatür", "Tümü", "Malzemedeki asbest oranı; ıslatma ve önceden söküm verimi", "EKSİKTİ", "Varsayım olarak eklendi (sarı). Belediyelerin asbest envanter raporlarında malzeme türü ve miktarı var: bilgi edinme konusu."),
        ("Basit değerler", "Standart Değerler B6:B20", "Değerler metin olarak girilmiş ('0,086', '%0,01', '100–200')", "Sorun", "Formülde kullanılamaz; aralıklar tek hücrede. Alt/Üst sayısal sütunlara ayrılmalı."),
        ("Basit değerler", "Standart Değerler B19", "Türkiye işyeri sınır değeri 100 lif/L (0,1 lif/cm³, 8 saat)", "Kullanılmadı", "Çalışanlar için mesleki sınır; mahalle sakinlerinin maruziyetini kıyaslamak riski küçük gösterir. Kıyas için olağan kent havası ve Kobe değerleri daha uygun."),
        ("Basit değerler", "Standart Değerler B11-B13", "Durgun gece ve gündüz karışım yüksekliği, durgun gece rüzgârı (genel değerler)", "Yerine yerel veri",
         "Yıkım gündüz yapıldığı için gece değerleri kullanılmadı. Kadıköy çalışma saatleri (ERA5 2015-2024): karışım yüksekliği medyanı 640 m (alt %10: 205 m); "
         "kararlı (E-F) sabah saatlerinde medyan 190 m, rüzgâr medyanı 2,8 m/s; rüzgâr <1 m/s saatlerin %2,9'u, <2 m/s %12,6'sı."),
    ]
    for i, r in enumerate(satirlar, 2):
        for j, v in enumerate(r, 1):
            c = yaz(ws, i, j, v, wrap=True)
            if r[3] in ("EKSİKTİ", "Sorun"):
                c.font = KALIN if j == 4 else SIYAH
    ws.freeze_panes = "A2"


# ----------------------------------------------------------------------------------------------------------------
def parametreler(wb, ws):
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 52
    for col, w in zip("CDEFGH", (12, 12, 12, 10, 70, 12)):
        ws.column_dimensions[col].width = w
    yaz(ws, 1, 1, "Parametreler — sarı ve mavi hücreler değiştirilebilir; Kaynak mahalleler ve Sonuçlar sayfaları buradan beslenir", KALIN)
    baslik_satiri(ws, 3, ["Kod", "Parametre", "Temel", "Alt", "Üst", "Birim", "Kaynak / gerekçe", "Tür"])
    # Kadıköy serisi tablosu (p ve M formülle buradan gelir)
    a = kadikoy_asbest()
    r0 = 22
    yaz(ws, r0 - 1, 1, "Kadıköy Açık Veri 'Asbest Verileri' (2026-04-06 güncellemesi) — p_asbest ve M_acm buradan formülle hesaplanır", KALIN)
    baslik_satiri(ws, r0, ["Yıl", "Denetlenen inşaat sahası", "Asbest tespit edilen bina", "Bertaraf (ton)", "Tespit oranı", "kg / tespit edilen bina"])
    for i, y in enumerate(a["yillar"]):
        r = r0 + 1 + i
        yaz(ws, r, 1, str(y))
        yaz(ws, r, 2, a["denetim"][i], MAVI)
        yaz(ws, r, 3, a["tespit"][i], MAVI)
        yaz(ws, r, 4, a["ton"][i], MAVI)
        yaz(ws, r, 5, f"=C{r}/B{r}", fmt="0.0%")
        yaz(ws, r, 6, f"=1000*D{r}/C{r}", fmt="#,##0")
    rs = r0 + 1 + len(a["yillar"])
    yaz(ws, rs, 1, "Havuz", KALIN)
    for col in "BCD":
        yaz(ws, rs, " BCD".index(col) + 1, f"=SUM({col}{r0+1}:{col}{rs-1})", KALIN)
    yaz(ws, rs, 5, f"=C{rs}/B{rs}", KALIN, fmt="0.0%")
    yaz(ws, rs, 6, f"=1000*D{rs}/C{rs}", KALIN, fmt="#,##0")
    yaz(ws, rs + 1, 1, "En düşük")
    yaz(ws, rs + 1, 5, f"=MIN(E{r0+1}:E{rs-1})", fmt="0.0%")
    yaz(ws, rs + 1, 6, f"=MIN(F{r0+1}:F{rs-1})", fmt="#,##0")
    yaz(ws, rs + 2, 1, "En yüksek")
    yaz(ws, rs + 2, 5, f"=MAX(E{r0+1}:E{rs-1})", fmt="0.0%")
    yaz(ws, rs + 2, 6, f"=MAX(F{r0+1}:F{rs-1})", fmt="#,##0")
    yaz(ws, rs + 3, 1, "Kaynak: https://acikveri.kadikoy.bel.tr/dataset/asbest-verileri (CSV). Payda 'denetlenen inşaat sahası', pay 'bina': oran yaklaşık.", GRI)

    satirlar = [
        ("hasar_carpani", P["hasar_carpani"], None),
        ("f_orta", P["f_orta"], None),
        ("p_asbest", P["p_asbest"], (f"=E{rs}", f"=E{rs+1}", f"=E{rs+2}")),
        ("M_acm", P["M_acm"], (f"=F{rs}", f"=F{rs+1}", f"=F{rs+2}")),
        ("w_asbest", P["w_asbest"], None),
        ("EF", P["EF"], None),
        ("k_lif", P["k_lif"], None),
        ("r_sokum", P["r_sokum"], None),
        ("eta_islatma", P["eta_islatma"], None),
        ("T_ay", P["T_ay"], None),
        ("kobe_duzeltme", P["kobe_duzeltme"], None),
    ]
    adres = {}
    for i, (kod, (b, lo, hi, birim, acik, kaynak, tur), formul) in enumerate(satirlar, 4):
        yaz(ws, i, 1, kod)
        yaz(ws, i, 2, acik, wrap=True)
        if formul:
            yaz(ws, i, 3, formul[0]); yaz(ws, i, 4, formul[1]); yaz(ws, i, 5, formul[2])
        else:
            kaynakli = kod in ("EF", "k_lif")
            yaz(ws, i, 3, b, MAVI if kaynakli else SIYAH, None if kaynakli else SARI)
            yaz(ws, i, 4, lo, SIYAH, SARI)
            yaz(ws, i, 5, hi, SIYAH, SARI)
        fmt = "0.0%" if birim == "oran" else ("0.00E+00" if kod == "EF" else "#,##0.##")
        for c in (3, 4, 5):
            ws.cell(row=i, column=c).number_format = fmt
        yaz(ws, i, 6, birim)
        yaz(ws, i, 7, kaynak, wrap=True)
        yaz(ws, i, 8, tur)
        adres[kod] = f"Parametreler!$C${i}"
        ad_tanimla(wb, f"par_{kod}", adres[kod])
    ek = len(satirlar) + 4
    sabitler = [
        ("calisma_saati_ay", "Ayda çalışma saati (08-18, pazartesi-cumartesi)", CALISMA_SAATI_AY, "saat", "Varsayım: yıkım ve enkaz işleri gündüz", True),
        ("calisma_payi", "Çalışma saatlerinin tüm saatlere oranı (ERA5 takvimi 2015-2024)", CALISMA_PAYI, "oran", "Dönem ortalamasına çevirmek için", False),
        ("arka_plan_alt", "Olağan kent havası, konut alanı (Kobe 1995, PCM)", S["arka_plan"][0], "lif/L", "Terazono ve ark. 2000; literatür tablonuz Ölçümler M04", False),
        ("arka_plan_ust", "Olağan kent havası, ana yol (Kobe 1995, PCM)", S["arka_plan"][1], "lif/L", "Terazono ve ark. 2000; literatür tablonuz Ölçümler M03", False),
    ]
    for j, (kod, acik, v, birim, kaynak, varsayim) in enumerate(sabitler):
        r = ek + j
        yaz(ws, r, 1, kod); yaz(ws, r, 2, acik, wrap=True)
        c = yaz(ws, r, 3, v, SIYAH if varsayim else MAVI, SARI if varsayim else None)
        c.number_format = "0.0%" if birim == "oran" else "#,##0.00"
        yaz(ws, r, 6, birim); yaz(ws, r, 7, kaynak, wrap=True)
        ad_tanimla(wb, f"par_{kod}", f"Parametreler!$C${r}")
    ws.freeze_panes = "A4"
    return rs


# ----------------------------------------------------------------------------------------------------------------
def kaynak_mahalleler(ws, kad, kaynak, kom, mah):
    md = {d["mahalle"]: d for d in mah}
    komd = {f"{k['ilce']}/{k['mahalle']}": k for k in kom if "hucre_sayisi" in k}
    basl = ["No", "İlçe", "Kaynak mahalle (D sayfalarındaki ad)", "Çok ağır hasarlı", "Ağır hasarlı", "Orta hasarlı",
            "Yarıçap içindeki pay", "Yıkılan bina", "Asbestli bina", "Asbest (kg)", "Havaya geçen asbest (kg)",
            "Lif sayısı", "Salım hızı Q (lif/s, çalışma saati)"]
    baslik_satiri(ws, 1, basl, [5, 11, 30, 10, 10, 10, 11, 11, 11, 12, 13, 14, 16])
    ws.cell(row=1, column=4).comment = Comment("İBB Açık Veri, Deprem Senaryosu Analiz Sonuçları (İBB/KOERİ 2019, Mw 7,5)", "model")
    for i, ad in enumerate(kaynak):
        r = i + 2
        if ad in md:
            d = md[ad]; ilce, mad = "Kadıköy", ad
            ca, ag, ot, pay = d["hasar_cok_agir"], d["hasar_agir"], d["hasar_orta"], 1.0
        else:
            k = komd[ad]; ilce, mad = k["ilce"], ad
            ca, ag, ot, pay = k["hasar_cok_agir"], k["hasar_agir"], k["hasar_orta"], k["pay_icerde"]
        yaz(ws, r, 1, i + 1); yaz(ws, r, 2, ilce); yaz(ws, r, 3, mad)
        yaz(ws, r, 4, ca, MAVI); yaz(ws, r, 5, ag, MAVI); yaz(ws, r, 6, ot, MAVI)
        yaz(ws, r, 7, float(pay), GRI, fmt="0%")
        yaz(ws, r, 8, f"=(D{r}+E{r}+par_f_orta*F{r})*par_hasar_carpani*G{r}", fmt="#,##0.0")
        yaz(ws, r, 9, f"=H{r}*par_p_asbest", fmt="#,##0.0")
        yaz(ws, r, 10, f"=I{r}*par_M_acm*par_w_asbest", fmt="#,##0")
        yaz(ws, r, 11, f"=J{r}*par_EF*(1-par_r_sokum)*(1-par_eta_islatma)*par_kobe_duzeltme", fmt="0.000")
        yaz(ws, r, 12, f"=K{r}*par_k_lif*1E12", fmt="0.00E+00")
        yaz(ws, r, 13, f"=L{r}/(par_T_ay*par_calisma_saati_ay*3600)", fmt="#,##0")
    n = len(kaynak)
    nK = len(kad)
    r = n + 3
    yaz(ws, r, 3, "Toplam Kadıköy", KALIN)
    for col in "DEFHIJKLM":
        c = yaz(ws, r, " ABCDEFGHIJKLM".index(col), f"=SUM({col}2:{col}{nK+1})", KALIN)
        c.number_format = ws[f"{col}2"].number_format
    yaz(ws, r + 1, 3, "Toplam komşu (2,5 km içi, alanla ölçekli)", KALIN)
    for col in "HIJKLM":
        c = yaz(ws, r + 1, " ABCDEFGHIJKLM".index(col), f"=SUM({col}{nK+2}:{col}{n+1})", KALIN)
        c.number_format = ws[f"{col}2"].number_format
    yaz(ws, r + 3, 1, "Yıkılan bina = (çok ağır + ağır + f_orta × orta) × hasar çarpanı × yarıçap içindeki pay. Komşu mahalle sınırları OSM (ODbL); "
                      "hasar sayıları İBB. Komşulara Kadıköy'ün asbest olasılığı ve bina başına malzeme miktarı uygulandı (varsayım).", GRI)
    ws.freeze_panes = "D2"


def d_sayfasi(wb, ad, D, kad, kaynak, aciklama):
    ws = wb.create_sheet(ad)
    n = len(kaynak)
    son = get_column_letter(1 + n)
    yaz(ws, 1, 1, "Alıcı ↓  /  Kaynak →", KALIN)
    yaz(ws, 2, 1, "Q (lif/s)", KALIN)
    for j, k in enumerate(kaynak):
        col = j + 2
        c = yaz(ws, 1, col, k, KALIN)
        c.alignment = Alignment(text_rotation=90)
        yaz(ws, 2, col, f"='Kaynak mahalleler'!M{j+2}", fmt="0.00E+00")
    yaz(ws, 3, 1, aciklama, GRI)
    sonuc_col = n + 3
    yaz(ws, 1, sonuc_col, "Derişim artışı (lif/L)", KALIN)
    for i, a in enumerate(kad):
        r = i + 4
        yaz(ws, r, 1, a, KALIN)
        for j in range(n):
            yaz(ws, r, j + 2, float(D[i, j]), GRI, fmt="0.00E+00")
        yaz(ws, r, sonuc_col, f"=SUMPRODUCT($B$2:${son}$2,B{r}:{son}{r})/1000", fmt="0.00000")
    ws.column_dimensions["A"].width = 16
    ws.row_dimensions[1].height = 150
    ws.freeze_panes = "B4"
    return get_column_letter(sonuc_col)


def sonuclar(ws, kad, d_kolon, mah):
    md = {d["mahalle"]: d for d in mah}
    sd = {d["mahalle"]: d for d in S["mahalle"]}
    basl = ["Mahalle", "Alan (km²)", "Nüfus 2025", "65 yaş üstü (2024)", "Geçici barınma ihtiyacı (İBB)", "Toplanma alanı (m²)",
            "Yıkılan bina (temel)", "Yıkım yoğunluğu (bina/km²)",
            "Derişim artışı, çalışma saatleri (lif/L)", "Dönem ortalaması (lif/L)", "Toplam maruziyet (lif/L·yıl)",
            "Olağan kent havasına oranı (alt sınır)", "Sıra", "Nüfus × derişim payı",
            "Kış", "Yaz", "Poyraz", "Lodos", "Kararlı sabah", "Durgun", "Kararlı sabah / uzun dönem",
            "MC %5", "MC %50", "MC %95", "MC: ilk 5'te olma olasılığı", "MC sıra aralığı (%5-%95)",
            "Kentsel dönüşüm dağılımında derişim", "Kentsel dönüşüm dağılımında sıra"]
    baslik_satiri(ws, 1, basl, [14, 8, 10, 10, 11, 11, 9, 10, 13, 12, 12, 12, 6, 10, 9, 9, 9, 9, 10, 9, 10, 9, 9, 9, 11, 12, 13, 11])
    n = len(kad)
    son = n + 1
    for i, a in enumerate(kad):
        r = i + 2
        d, s = md[a], sd[a]
        yaz(ws, r, 1, a, KALIN)
        yaz(ws, r, 2, round(d["alan_km2"], 3), GRI)
        yaz(ws, r, 3, d["nufus_2025"], MAVI, fmt="#,##0")
        yaz(ws, r, 4, d["nufus_65_ustu_2024"], MAVI, fmt="#,##0")
        yaz(ws, r, 5, d["gecici_barinma"], MAVI, fmt="#,##0")
        yaz(ws, r, 6, d["toplanma_alani_m2"], MAVI, fmt="#,##0")
        yaz(ws, r, 7, f"='Kaynak mahalleler'!H{i+2}", fmt="#,##0.0")
        yaz(ws, r, 8, f"=G{r}/B{r}", fmt="#,##0.0")
        yaz(ws, r, 9, f"='D_Temel'!{d_kolon['D_Temel']}{r+2}", fmt="0.0000")
        yaz(ws, r, 10, f"=I{r}*par_calisma_payi", fmt="0.0000")
        yaz(ws, r, 11, f"=J{r}*par_T_ay/12", fmt="0.0000")
        yaz(ws, r, 12, f"=I{r}/par_arka_plan_alt", fmt="0.0%")
        yaz(ws, r, 13, f"=RANK(I{r},$I$2:$I${son})")
        yaz(ws, r, 14, f"=C{r}*I{r}/SUMPRODUCT($C$2:$C${son},$I$2:$I${son})", fmt="0.0%")
        for col, sh in zip(range(15, 21), ("D_Kış", "D_Yaz", "D_Poyraz", "D_Lodos", "D_Kararlı", "D_Durgun")):
            yaz(ws, r, col, f"='{sh}'!{d_kolon[sh]}{r+2}", fmt="0.0000")
        yaz(ws, r, 21, f"=S{r}/I{r}", fmt="0.0")
        yaz(ws, r, 22, s["mc_p5"], GRI, fmt="0.0000")
        yaz(ws, r, 23, s["mc_p50"], GRI, fmt="0.0000")
        yaz(ws, r, 24, s["mc_p95"], GRI, fmt="0.000")
        yaz(ws, r, 25, s["mc_ilk5"], GRI, fmt="0%")
        yaz(ws, r, 26, f"{s['mc_sira_p5']:.0f}-{s['mc_sira_p95']:.0f}", GRI)
        yaz(ws, r, 27, s["c_riskli_yapi"], GRI, fmt="0.0000")
        yaz(ws, r, 28, s["sira_riskli_yapi"], GRI)
    r = son + 2
    yaz(ws, r, 1, "İlçe özeti", KALIN)
    ozet = [
        ("Nüfus ağırlıklı ortalama, çalışma saatleri (lif/L)", f"=SUMPRODUCT(C2:C{son},I2:I{son})/SUM(C2:C{son})", "0.0000"),
        ("Nüfus ağırlıklı ortalama, dönem (lif/L)", f"=SUMPRODUCT(C2:C{son},J2:J{son})/SUM(C2:C{son})", "0.0000"),
        ("En yüksek mahalle derişimi (lif/L)", f"=MAX(I2:I{son})", "0.0000"),
        ("En yüksek mahalle", f"=INDEX(A2:A{son},MATCH(MAX(I2:I{son}),I2:I{son},0))", None),
        ("En yüksek / olağan kent havası (alt)", f"=MAX(I2:I{son})/par_arka_plan_alt", "0.0%"),
        ("Kararlı sabah en yüksek (lif/L)", f"=MAX(S2:S{son})", "0.0000"),
    ]
    for j, (et, f, fmt) in enumerate(ozet, 1):
        yaz(ws, r + j, 1, et)
        yaz(ws, r + j, 9, f, KALIN, fmt=fmt)
    r = r + len(ozet) + 2
    notlar = [
        "Derişimler yıkımın eklediği artıştır (arka plan hariç), PCM eşdeğeri lif varsayımıyla. 'Çalışma saatleri' = yıkım yapılan saatlerdeki ortalama; 'dönem ortalaması' = gece dahil.",
        "Toplam maruziyet = dönem ortalaması × yıkım süresi: süreden bağımsızdır (aynı lif daha kısa sürede salınırsa ortalama artar, toplam değişmez).",
        "Kış/Yaz: yalnız o mevsimin çalışma saatleri. Poyraz/Lodos/Kararlı sabah/Durgun: o hava yıkım süresince sürseydi (kısa süreli bölümlerin göstergesi). Tanımlar Meteoroloji sayfasında.",
        "MC: Monte Carlo (4.000 deneme; İBB hasar dağılımı, mahalle hasar sayısında log-normal belirsizlik). Kentsel dönüşüm dağılımı: yıkımlar 2012-2022 riskli yapı tespitleri oranında. Gri sütunlar Python çıktısı.",
    ]
    for j, t in enumerate(notlar):
        yaz(ws, r + j, 1, t, GRI)
    ws.freeze_panes = "B2"


def meteoroloji(ws):
    yaz(ws, 1, 1, "ERA5 saatlik (41,0K 29,0D hücresi, Open-Meteo Historical Weather API), 2015-2024; Pasquill-Gifford kararlılığı Turner yöntemiyle", KALIN)
    sutunlar = [("tum_saatler", "Tüm saatler"), ("calisma_saatleri", "Çalışma saatleri (08-18, pzt-cmt)"),
                ("kis_calisma", "Kış çalışma"), ("yaz_calisma", "Yaz çalışma"), ("sabah_08_10", "Sabah 08-10")]
    baslik_satiri(ws, 3, ["Rüzgârın geldiği yön (%)"] + [b for _, b in sutunlar], [26, 14, 18, 14, 14, 14])
    yonler = list(MO["tum_saatler"]["gul"].keys())
    for i, y in enumerate(yonler, 4):
        yaz(ws, i, 1, y)
        for j, (k, _) in enumerate(sutunlar, 2):
            yaz(ws, i, j, MO[k]["gul"][y], MAVI, fmt="0.0%")
    r = 4 + len(yonler) + 1
    baslik_satiri(ws, r, ["Kararlılık sınıfı (%)"] + [b for _, b in sutunlar])
    for i, sn in enumerate("ABCDEF", r + 1):
        yaz(ws, i, 1, {"A": "A çok kararsız", "B": "B kararsız", "C": "C hafif kararsız", "D": "D nötr", "E": "E hafif kararlı", "F": "F kararlı"}[sn])
        for j, (k, _) in enumerate(sutunlar, 2):
            yaz(ws, i, j, MO[k]["sinif"][sn], MAVI, fmt="0.0%")
    r = r + 8
    baslik_satiri(ws, r, ["Özet"] + [b for _, b in sutunlar])
    olcu = [("saat", "Saat sayısı", "#,##0"), ("ruzgar_ortalama", "Ortalama rüzgâr (m/s)", "0.00"),
            ("sakin_payi_<1ms", "Rüzgâr < 1 m/s", "0.0%"), ("hafif_payi_<2ms", "Rüzgâr < 2 m/s", "0.0%"),
            ("yagisli_saat_payi_>0.1mm", "Yağışlı saat (>0,1 mm)", "0.0%"), ("blh_medyan", "Karışım yüksekliği medyanı (m)", "#,##0"),
            ("blh_p10", "Karışım yüksekliği alt %10 (m)", "#,##0")]
    for i, (k, et, fmt) in enumerate(olcu, r + 1):
        yaz(ws, i, 1, et)
        for j, (kk, _) in enumerate(sutunlar, 2):
            yaz(ws, i, j, MO[kk][k], MAVI, fmt=fmt)
    r = r + len(olcu) + 2
    yaz(ws, r, 1, "Kısa süreli senaryolar (D_Poyraz, D_Lodos, D_Kararlı, D_Durgun)", KALIN)
    sen = [
        ("Poyraz", S["senaryolar"]["S1_poyraz"], "Çalışma saatlerinin %56'sında rüzgâr K-DKD arasından (en sık KD, %26). Bu saatlerde rüzgâr medyanı 5,3 m/s (KD'de 5,8), karışım yüksekliği medyanı 800 m (KD'de 875)."),
        ("Lodos", S["senaryolar"]["S2_lodos"], "Çalışma saatlerinin %22'sinde GGB-BGB (kışın %30). Senaryo değerleri bu saatlerin medyanları: 3,9 m/s, 485 m."),
        ("Kararlı sabah", S["senaryolar"]["S3_kararli_sabah"], "Çalışma saatlerinin %1,6'sı E-F; hepsi 08-10 arasında (sabahların %7,9'u). Bu saatlerin %55'inde rüzgâr güney yarısından (en sık G, %15,5); "
                                                              "rüzgâr medyanı 2,8 m/s (alt %10: 1,4), karışım yüksekliği medyanı 190 m (alt %25: 105). Senaryo düşük rüzgârlı ucu temsil eder."),
        ("Durgun", S["senaryolar"]["S4_durgun"], "Çalışma saatlerinin %2,9'unda rüzgâr < 1 m/s; karışım yüksekliği medyanı 255 m. Bu saatlerin çoğu güneşli öğle (A-B); "
                                                 "senaryoda nötr (D) varsayıldı, yani üst sınıra yakın. ERA5 durgunluğu az gösterebilir."),
    ]
    for i, (a, t, g) in enumerate(sen, r + 1):
        yaz(ws, i, 1, a, KALIN); yaz(ws, i, 2, t); yaz(ws, i, 4, g, GRI)
    r = r + len(sen) + 2
    yaz(ws, r, 1, "Not: kıyıdaki 0,25° ERA5 hücresi kara ve deniz karışımıdır; durgun ve kararlı saatleri istasyon ölçümünden az gösterebilir. MGM Göztepe istasyonu saatlik verisiyle karşılaştırılmalı.", GRI)


def duyarlilik(ws):
    yaz(ws, 1, 1, "Duyarlılık analizi (Python çıktısı; Excel'de yeniden hesaplanmaz)", KALIN)
    yaz(ws, 2, 1, "A) Tek tek değiştirme: her parametre aralığının ucuna çekildiğinde", KALIN)
    baslik_satiri(ws, 3, ["Parametre", "Açıklama", "Tür", "Uç", "Değer", "Temel değer", "İlçe ortalaması çarpanı",
                          "En yüksek mahalle çarpanı", "Sıralama benzerliği (Spearman ρ)", "İlk 5'ten kalan"],
                  [14, 50, 11, 6, 11, 11, 12, 12, 13, 10])
    for i, o in enumerate(sorted(DU["oat"], key=lambda o: -abs(np.log(o["pwm_carpan"]))), 4):
        for j, v in enumerate([o["parametre"], o["aciklama"], o["tur"], o["uc"], o["deger"], o["temel_deger"],
                               o["pwm_carpan"], o["maks_carpan"], o["rho"], o["ilk5_kalan"]], 1):
            c = yaz(ws, i, j, v, GRI)
            if j in (5, 6):
                c.number_format = "0.00E+00" if o["parametre"] == "EF" else "#,##0.###"
            if j in (7, 8):
                c.number_format = "0.00"
            if j == 9:
                c.number_format = "0.000"
    r = 4 + len(DU["oat"]) + 1
    yaz(ws, r, 1, "B) Yapısal seçenekler", KALIN)
    baslik_satiri(ws, r + 1, ["Seçenek", "", "", "", "", "", "İlçe ortalaması çarpanı", "En yüksek mahalle çarpanı",
                              "Sıralama benzerliği (Spearman ρ)", "İlk 5'ten kalan", "İlk 5"])
    ws.column_dimensions["K"].width = 60
    for i, y in enumerate(DU["yapisal"], r + 2):
        yaz(ws, i, 1, y["secenek"], GRI)
        for j, k in ((7, "pwm_carpan"), (8, "maks_carpan"), (9, "rho"), (10, "ilk5_kalan")):
            c = yaz(ws, i, j, y[k], GRI)
            c.number_format = "0.000" if k == "rho" else "0.00" if k != "ilk5_kalan" else "0"
        yaz(ws, i, 11, y["ilk5"], GRI)
    r = r + 2 + len(DU["yapisal"]) + 1
    yaz(ws, r, 1, "C) Monte Carlo (4.000 deneme): log(ilçe ortalaması) varyansının girdilere göre payı", KALIN)
    yaz(ws, r, 7, f"Log-doğrusal R² = {DU['mc']['r2_logdogrusal']:.3f}", GRI)
    baslik_satiri(ws, r + 1, ["Girdi", "", "", "", "", "", "Varyans payı", "Std. regresyon katsayısı"])
    for i, v in enumerate(DU["mc"]["varyans"], r + 2):
        yaz(ws, i, 1, v["girdi"], GRI)
        yaz(ws, i, 7, v["pay"], GRI, fmt="0.0%")
        yaz(ws, i, 8, v["src"], GRI, fmt="0.000")
    r = r + 2 + len(DU["mc"]["varyans"]) + 1
    yaz(ws, r, 1, "D) Monte Carlo yüzdelikleri", KALIN)
    baslik_satiri(ws, r + 1, ["Çıktı", "", "", "", "", "", "%5", "%25", "%50", "%75", "%95"])
    for i, (et, k) in enumerate((("İlçe nüfus ağırlıklı ortalama, çalışma saatleri (lif/L)", "pwm_yuzdelik"),
                                 ("Toplam maruziyet, nüfus ağırlıklı (lif/L·yıl)", "kumulatif_yuzdelik")), r + 2):
        yaz(ws, i, 1, et, GRI)
        for j, q in enumerate(("5", "25", "50", "75", "95"), 7):
            yaz(ws, i, j, DU["mc"][k][q], GRI, fmt="0.0000")
    r = r + 5
    yaz(ws, r, 1, "Örnekleme: hasar çarpanı log-tekdüze 0,5-2; f_orta tekdüze 0-1; p_asbest ve w üçgen; M_acm log-üçgen; EF log-tekdüze 1E-5–1E-3; k_lif log-tekdüze 10-100; "
                  "söküm ve ıslatma üçgen (en olası 0); süre tekdüze 6-24 ay; Kobe düzeltmesi log-tekdüze 1-5; yayılım seçeneği (temel/kentsel/2 m/15 m) ve asbest olasılığının mekânsal dağılımı rastgele; "
                  "mahalle hasar sayısında log-normal σ=0,5.", GRI)


def tek_yikim(ws):
    yaz(ws, 1, 1, "Asbestli tek bir binanın kontrolsüz yıkımı: rüzgâr altı eksen derişimi (Gauss dumanı, Briggs kentsel)", KALIN)
    girdi = [
        ("Bina başına asbestli malzeme (kg)", "=par_M_acm", None), ("Malzemedeki asbest oranı", "=par_w_asbest", "0.0%"),
        ("Havaya geçen pay (EF)", "=par_EF", "0.00E+00"), ("Islatma verimi", "=par_eta_islatma", "0%"),
        ("Lif dönüşümü (lif/ng)", "=par_k_lif", None), ("Yıkım süresi (saat)", 40, None),
        ("Salım yüksekliği h (m)", 5, None), ("Ölçüm yüksekliği z (m)", 1.5, None),
        ("Başlangıç σy0 (m)", 5, None), ("Başlangıç σz0 (m)", 7, None),
    ]
    for i, (et, v, fmt) in enumerate(girdi, 3):
        yaz(ws, i, 1, et)
        c = yaz(ws, i, 2, v, SIYAH, SARI if not isinstance(v, str) else None)
        if fmt:
            c.number_format = fmt
    yaz(ws, 13, 1, "Salım hızı Q (lif/s)", KALIN)
    yaz(ws, 13, 2, "=B3*B4*B5*(1-B6)*B7*1E12/(B8*3600)", KALIN, fmt="0.00E+00")
    # Briggs kentsel katsayıları
    yaz(ws, 15, 1, "Briggs kentsel katsayıları: σy = cy·x·(1+0,0004x)^-0,5 ; σz = cz·x·(1+dz·x)^ez", KALIN)
    baslik_satiri(ws, 16, ["Sınıf", "cy", "cz", "dz", "ez"])
    kats = [("AB", 0.32, 0.24, 0.001, 0.5), ("C", 0.22, 0.20, 0.0, 0.0), ("D", 0.16, 0.14, 0.0003, -0.5), ("EF", 0.11, 0.08, 0.0015, -0.5)]
    for i, k in enumerate(kats, 17):
        for j, v in enumerate(k, 1):
            yaz(ws, i, j, v, MAVI if j > 1 else SIYAH)
    yaz(ws, 21, 1, "Kaynak: Briggs (1973), McElroy-Pooler St. Louis deneylerine dayalı kentsel eğriler (100 m-10 km).", GRI)
    mesafe = [20, 50, 100, 200, 500, 1000]
    baslik_satiri(ws, 23, ["Hava durumu", "Rüzgâr u (m/s)", "Sınıf"] + [f"{m} m (lif/L)" for m in mesafe])
    for j, m in enumerate(mesafe, 4):
        yaz(ws, 22, j, m, KALIN)
    havalar = [("Poyraz (KD)", 5.5, "D"), ("Lodos (GB)", 3.9, "C"), ("Kararlı sabah", 1.5, "EF"), ("Durgun", 1.0, "D")]
    yaz(ws, 22, 3, "Mesafe (m) →", KALIN)
    durumlar = [("Temel salım (ıslatmasız yıkım)", 1.0),
                ("Gevrek malzeme / kuru kırma (EF ×10)", 10.0),
                ("Islatmalı yıkım (liflerin %70'i tutulur)", 0.3)]
    i = 24
    for dad, carpan in durumlar:
        yaz(ws, i, 1, dad, KALIN)
        c = yaz(ws, i, 3, carpan, SIYAH, SARI)
        c.number_format = '"×"0.0'
        c.comment = Comment("Salım çarpanı: yukarıdaki Q ile çarpılır", "model")
        carp = f"$C${i}"
        for a, u, sn in havalar:
            i += 1
            yaz(ws, i, 1, a); yaz(ws, i, 2, u, SIYAH, SARI); yaz(ws, i, 3, sn, SIYAH, SARI)
            for j in range(len(mesafe)):
                col = get_column_letter(4 + j)
                x = f"{col}$22"
                cy = f"VLOOKUP($C{i},$A$17:$E$20,2,FALSE)"
                cz = f"VLOOKUP($C{i},$A$17:$E$20,3,FALSE)"
                dz = f"VLOOKUP($C{i},$A$17:$E$20,4,FALSE)"
                ez = f"VLOOKUP($C{i},$A$17:$E$20,5,FALSE)"
                sy = f"SQRT(({cy}*{x}*(1+0.0004*{x})^-0.5)^2+$B$11^2)"
                sz = f"SQRT(({cz}*{x}*(1+{dz}*{x})^{ez})^2+$B$12^2)"
                # Excel'de tekli eksi üs almadan önce uygulanır (-a^2 = +a²); bu yüzden -0,5*(...)^2 biçimi
                v = f"(EXP(-0.5*(($B$10-$B$9)/{sz})^2)+EXP(-0.5*(($B$10+$B$9)/{sz})^2))"
                yaz(ws, i, 4 + j, f"=$B$13*{carp}*{v}/(2*PI()*{sy}*{sz}*$B{i})/1000", fmt="0.000")
        i += 2
    yaz(ws, i, 1, "Karşılaştırma: olağan kent havası 0,19-0,41 lif/L; Kobe deprem sonrası kent havası 1-2 lif/L; Kobe'de sökümsüz yıkımın parsel sınırı 160-250 lif/L (PCM); "
                  "Malatya 2024 aktif yıkım alanı 4-10 lif/L (PCM lifli toz, asbest teyidi yok).", GRI)
    yaz(ws, i + 1, 1, "20-50 m için Briggs eğrileri dış değerlemedir (100 m altı ölçülmemiş). Yakın alanda Gauss modelleri ölçümü 1,6-2,6 kat düşük tahmin edebilir (Duke Forest 1996). "
                      "Karışım tavanı yansıması burada yok; Python sürümünde var, fark %1'in altında.", GRI)
    ws.column_dimensions["A"].width = 34
    for j in range(2, 10):
        ws.column_dimensions[get_column_letter(j)].width = 13


def kaynaklar(ws):
    baslik_satiri(ws, 1, ["Veri / kaynak", "Ne için", "Bağlantı", "Erişim (UTC)", "Lisans / not"], [44, 40, 70, 22, 50])
    ibb = json.load(open(RAW / "ibb" / "deprem_senaryosu_analiz_sonuclari.meta.json"))
    era = json.load(open(RAW / "meteo" / "era5_kadikoy_2015_2024.meta.json"))
    kad = {d["dosya"]: d for d in json.load(open(RAW / "kadikoy" / "model" / "indirilenler.json"))}
    satirlar = [
        ("İBB Açık Veri — Deprem Senaryosu Analiz Sonuçları (İBB/KOERİ 2019, Mw 7,5)", "Mahalle bazında hasarlı bina, geçici barınma", ibb["url"], ibb["erisim"], f"İBB Açık Veri Lisansı; SHA-256 {ibb['sha256'][:16]}…"),
        ("Kadıköy Açık Veri — Asbest Verileri 2015-2025", "Asbest olasılığı, bina başına malzeme", "https://acikveri.kadikoy.bel.tr/dataset/asbest-verileri", "2026-10-13", "Kadıköy Belediyesi açık veri"),
        ("Kadıköy Açık Veri — Riskli Yapı Tespiti, mahalle 2012-2022", "Kentsel dönüşüm dağılımı (duyarlılık)", "https://acikveri.kadikoy.bel.tr/dataset/riskli-yapi-tespiti-yapilan-binalar", "2026-10-13", ""),
    ]
    for f, ne in (("kadikoy-ilcesindeki-mahalle-sinirlari.shp", "Mahalle sınırları (TM30)"), ("mahalle-nufuslarinin-cinsiyete-gore-dagilimi.csv", "2025 nüfus"),
                  ("65-uzeri-nufusun-mahalle-dagilimlari.csv", "2024 65 yaş üstü"), ("yapim-yillarina-gore-binalarin-mekansal-dagilimi.csv", "Bina yapım yılı 2021"),
                  ("mahalle-afet-toplanma-alanlari-miktari.csv", "Afet toplanma alanları"), ("kat-sayisina-gore-binalarin-mekansal-dagilimi.csv", "Kat sayısı 2021"),
                  ("insaat-turlerine-gore-binalarin-mekansal-dagilimi.csv", "İnşaat türü 2021"), ("mahalle-bazinda-konut-isyeri-dagilimi.csv", "Konut-işyeri 2022")):
        d = kad[f]
        satirlar.append((f"Kadıköy Açık Veri — {f}", ne, d["url"], d["erisim"], f"SHA-256 {d['sha256'][:16]}…"))
    satirlar += [
        ("ERA5 saatlik (Open-Meteo Historical Weather API)", "Rüzgâr, ışınım, bulut, yağış, karışım yüksekliği 2015-2024", era["url"], era["erisim"], era["lisans"] + ". Not: open-meteo.com robots.txt tüm yolları kapatıyor; belgelenmiş ücretsiz API olduğu için tek istek istisnası yapıldı (README)."),
        ("OSM mahalle sınırları (sahircansurmeli/istanbul-geojson, commit c9bf151)", "Kadıköy poligonlarına ad vermek (IoU ≥ 0,945) ve komşu mahalleler", "https://github.com/sahircansurmeli/istanbul-geojson", "2026-10-13", "© OpenStreetMap katkıcıları, ODbL 1.0"),
        ("Terazono A., Sakai S., Takatsuki H. (2000). The Great Hanshin-Awaji Earthquake of Japan 1995 and asbestos emission. Air Pollution VIII, WIT Press", "Salım faktörü, lif dönüşümü, Kobe ölçümleri", "https://www.witpress.com/Secure/elibrary/papers/AIR00/AIR00057FU.pdf", "", "Sizin literatür tablonuz, Kaynak No 1"),
        ("Gilliam R.C., Huber A.H., Raman S. (2005). WTC CALPUFF", "Meteoroloji–seyrelme ilişkisi (niteliksel)", "https://sraman.wordpress.ncsu.edu/files/2020/09/Raman.Publication.Number.189.pdf", "", "Kaynak No 2"),
        ("Duke Forest (1996), OSTI 671861", "Yakın alan model belirsizliği 1,6-2,6 kat", "https://www.osti.gov/biblio/671861", "", "Kaynak No 18"),
        ("Briggs G.A. (1973). Diffusion estimation for small emissions. ATDL-106 (kentsel σ eğrileri)", "Yayılım katsayıları", "", "", "Standart düzenleyici eğriler (ISC/ISCLT)"),
        ("Turner D.B. (1970). Workbook of Atmospheric Dispersion Estimates (kararlılık sınıflaması)", "Pasquill-Gifford sınıfı", "", "", ""),
    ]
    for i, r in enumerate(satirlar, 2):
        for j, v in enumerate(r, 1):
            c = yaz(ws, i, j, v, wrap=True)
            if j == 3 and str(v).startswith("http"):
                c.hyperlink = v
                c.font = Font(name=F, size=10, color="0563C1", underline="single")


def haritalar(ws):
    yaz(ws, 1, 1, "Haritalar ve grafikler (model/cikti/*.png)", KALIN)
    yer = [("harita_temel.png", "A3", 0.42), ("harita_dagilim_karsilastirma.png", "L3", 0.42),
           ("harita_meteoroloji.png", "A33", 0.42), ("tornado.png", "A52", 0.42), ("yakin_alan.png", "L52", 0.42)]
    for f, hucre, olcek in yer:
        img = XLImage(str(CIKTI / f))
        img.width, img.height = img.width * olcek, img.height * olcek
        ws.add_image(img, hucre)


def main():
    D, kad, kaynak, kom, senaryolar, iz = aktarimlar()
    mah = mahalle_tablosu()
    wb = Workbook()
    ws = wb.active
    ws.title = "Okuma notu"
    okuma_notu(ws)
    tablolariniz(wb.create_sheet("Tablolarınız"))
    parametreler(wb, wb.create_sheet("Parametreler"))
    kaynak_mahalleler(wb.create_sheet("Kaynak mahalleler"), kad, kaynak, kom, mah)
    ws_son = wb.create_sheet("Sonuçlar")
    meteoroloji(wb.create_sheet("Meteoroloji"))
    duyarlilik(wb.create_sheet("Duyarlılık"))
    tek_yikim(wb.create_sheet("Tek yıkım"))
    haritalar(wb.create_sheet("Haritalar"))
    kaynaklar(wb.create_sheet("Kaynaklar"))
    d_kolon = {}
    for ad, k, ack in (("D_Temel", "temel", "Uzun dönem: tüm yıl çalışma saatleri (ERA5 2015-2024)"), ("D_Kış", "kis", "Aralık-Şubat çalışma saatleri"),
                       ("D_Yaz", "yaz", "Haziran-Ağustos çalışma saatleri"), ("D_Poyraz", "S1_poyraz", S["senaryolar"]["S1_poyraz"]),
                       ("D_Lodos", "S2_lodos", S["senaryolar"]["S2_lodos"]), ("D_Kararlı", "S3_kararli_sabah", S["senaryolar"]["S3_kararli_sabah"]),
                       ("D_Durgun", "S4_durgun", S["senaryolar"]["S4_durgun"])):
        d_kolon[ad] = d_sayfasi(wb, ad, D[k], kad, kaynak,
                                f"Aktarım katsayısı D (lif/m³ başına lif/s), Python yayılım modeli; {ack}. Sağdaki sütun = Σ Q×D / 1000.")
    sonuclar(ws_son, kad, d_kolon, mah)
    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for c in row:
                if c.value is not None and (c.font is None or c.font.name != F):
                    c.font = Font(name=F, size=c.font.size or 10, bold=c.font.bold, color=c.font.color)
    out = CIKTI / "kadikoy_asbest_yayilim_modeli.xlsx"
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
